"""Frozen native addresses are checked from exact originals, without caches."""

import importlib.util
import io
from pathlib import Path
import sys

import pytest
import yaml


DOMAIN = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location(
    "native_document_contracts", DOMAIN / "governance/contracts/__init__.py",
    submodule_search_locations=[str(DOMAIN / "governance/contracts")],
)
CONTRACTS = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = CONTRACTS
SPEC.loader.exec_module(CONTRACTS)
from native_document_contracts import checks, source_resources as sources, validation as v  # noqa: E402

FIXTURE_SPEC = importlib.util.spec_from_file_location(
    "native_document_fixture", Path(__file__).with_name("test_clean_clone.py")
)
FIXTURE = importlib.util.module_from_spec(FIXTURE_SPEC)
FIXTURE_SPEC.loader.exec_module(FIXTURE)


def frozen_source(kind):
    data, locator = FIXTURE.document_bytes(kind)
    text, addresses, extractor, _ = sources.extract_document(data, "." + kind)
    snapshot = {
        "source_id": "document", "snapshot_id": "native-document", "format": "." + kind,
        "sha256": v.sha256_bytes(data), "bytes": len(data),
        "extraction": {"extractor": extractor, "text_sha256": v.sha256_bytes(text.encode()),
                       "locator_sha256": v.sha256_bytes(v.canonical_bytes({"locators": addresses}))},
    }
    lock = {"resources": [{"source_id": "document", "kind": "raw"}], "snapshots": [snapshot]}
    ref = {"evidence_ref": "original", "source_id": "document", "snapshot_id": "native-document", "locator": locator}
    return data, snapshot, lock, ref


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_fixed_checks_replay_frozen_native_address_from_original_bytes(kind):
    data, _, lock, ref = frozen_source(kind)
    assert sources.validate_reference(ref, lock, lambda _: data)["locator"] == ref["locator"]


@pytest.mark.parametrize("kind", ["pdf", "docx"])
@pytest.mark.parametrize("field", ["text_sha256", "locator_sha256"])
def test_frozen_document_extraction_mismatch_is_never_accepted(kind, field):
    data, snapshot, lock, ref = frozen_source(kind)
    snapshot["extraction"][field] = "0" * 64
    with pytest.raises(v.ContractError) as error:
        sources.validate_reference(ref, lock, lambda _: data)
    assert error.value.code == "SOURCE_EXTRACTION_UNAVAILABLE"


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_document_adapter_version_mismatch_is_an_input_gap(kind):
    data, snapshot, lock, ref = frozen_source(kind)
    snapshot["extraction"]["extractor"] = "different-adapter:0"
    with pytest.raises(v.ContractError) as error:
        sources.validate_reference(ref, lock, lambda _: data)
    assert error.value.code == "SOURCE_ADAPTER_UNAVAILABLE"


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_native_checks_do_not_accept_changed_original_bytes(kind):
    data, _, lock, ref = frozen_source(kind)
    with pytest.raises(v.ContractError) as error:
        sources.validate_reference(ref, lock, lambda _: data + b"changed")
    assert error.value.code == "EVIDENCE_OBJECT_CHANGED"


def test_docx_native_table_rows_and_paragraph_ranges_use_real_addresses():
    import docx
    document = docx.Document()
    document.add_paragraph("")
    document.add_paragraph("First nonempty paragraph")
    document.add_paragraph("Second nonempty paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Native table cell"
    stream = io.BytesIO()
    document.save(stream)
    _, addresses, _, _ = sources.extract_document(stream.getvalue(), ".docx")
    selected = sources.document_segments(addresses, {"kind": "paragraphs", "start": 2, "end": 3}, ".docx")
    assert [row["paragraph"] for row in selected] == [2, 3]
    assert sources.document_segments(addresses, {"kind": "table_row", "table": 1, "row": 1}, ".docx")[0]["text"] == "Native table cell"
    data = stream.getvalue()
    text, addresses, extractor, _ = sources.extract_document(data, ".docx")
    snapshot = {"source_id": "document", "snapshot_id": "native-table", "format": ".docx",
                "sha256": v.sha256_bytes(data), "bytes": len(data),
                "extraction": {"extractor": extractor, "text_sha256": v.sha256_bytes(text.encode()),
                               "locator_sha256": v.sha256_bytes(v.canonical_bytes({"locators": addresses}))}}
    lock = {"resources": [{"source_id": "document", "kind": "raw"}], "snapshots": [snapshot]}
    ref = {"evidence_ref": "original", "source_id": "document", "snapshot_id": "native-table",
           "locator": {"kind": "table_row", "table": 1, "row": 1}}
    assert sources.validate_reference(ref, lock, lambda _: data)["locator"] == ref["locator"]
    for locator in ({"kind": "paragraph", "paragraph": 1},
                    {"kind": "table_row", "table": True, "row": 1},
                    {"kind": "table_row", "table": 1, "row": 2}):
        with pytest.raises(v.ContractError) as error:
            sources.document_segments(addresses, locator, ".docx")
        assert error.value.code == "INVALID_SOURCE_LOCATOR"


def test_native_page_ranges_reject_missing_pages_and_booleans():
    data, _, _, _ = frozen_source("pdf")
    _, addresses, _, _ = sources.extract_document(data, ".pdf")
    assert len(sources.document_segments(addresses, {"kind": "pages", "start": 1, "end": 1}, ".pdf")) == 1
    for locator in ({"kind": "page", "page": True}, {"kind": "pages", "start": 1, "end": 2}):
        with pytest.raises(v.ContractError):
            sources.document_segments(addresses, locator, ".pdf")


def test_code_file_named_pdf_keeps_its_frozen_code_adapter():
    data = b"int native_code = 1;\n"
    text, addresses = sources.extract_code_text(data, "code-utf8:1")
    snapshot = {"source_id": "code", "snapshot_id": "code-pdf-name", "format": ".pdf",
                "sha256": v.sha256_bytes(data), "bytes": len(data),
                "extraction": {"extractor": "code-utf8:1", "text_sha256": v.sha256_bytes(text.encode()),
                               "locator_sha256": v.sha256_bytes(v.canonical_bytes({"locators": addresses}))}}
    lock = {"resources": [{"source_id": "code", "kind": "code"}], "snapshots": [snapshot]}
    ref = {"evidence_ref": "code", "source_id": "code", "snapshot_id": "code-pdf-name",
           "locator": {"kind": "lines", "start": 1, "end": 1}}
    assert sources.validate_reference(ref, lock, lambda _: data)["snapshot_id"] == "code-pdf-name"


def test_legacy_explicit_structured_source_without_frozen_adapter_still_works():
    data, snapshot, lock, ref = frozen_source("pdf")
    snapshot.pop("extraction")
    ref["locator"] = {"kind": "page", "number": 1}
    explicit = {"bytes": data, "extracted_from_sha256": snapshot["sha256"], "locators": {"page": [1]}}
    assert sources.validate_reference(ref, lock, lambda _: explicit)["snapshot_id"] == snapshot["snapshot_id"]


def pipeline(tmp_path):
    result = FIXTURE.MANUAL.ManualPipeline(tmp_path)
    lock_path = result.root / "governance/framework.lock.yaml"
    lock = yaml.safe_load(lock_path.read_bytes())
    for entry in lock["files"].values():
        entry["sha256"] = v.sha256_bytes((result.root / entry["path"]).read_bytes())
    lock_path.write_text(yaml.safe_dump(lock, sort_keys=False), encoding="utf-8")
    return result


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_source_read_rejects_adapter_drift_when_extraction_cache_is_absent(tmp_path, kind):
    p = pipeline(tmp_path)
    data, locator = FIXTURE.document_bytes(kind)
    source_path = "docs-raw/specs/native." + kind
    p.write(source_path, data)
    bundle_ref = p.call("knowledge-source-capture", {"operation_id": "adapter-drift",
                       "source_path": source_path, "source_id": "document"})["data"]["bundle_ref"]
    bundle = p.record(bundle_ref)
    extraction = bundle["snapshots"][0]["extraction"]
    (p.root / extraction["text_ref"]).unlink()
    (p.root / extraction["locator_ref"]).unlink()
    extraction["extractor"] = "different-adapter:0"
    bundle["digest"] = v.digest_payload(bundle)
    p.write(bundle_ref, v.canonical_bytes(bundle))
    ref = {"evidence_ref": "original", "source_id": "document",
           "snapshot_id": bundle["snapshots"][0]["snapshot_id"], "locator": locator}
    response = p.call("knowledge-source-capture", {"operation_id": "adapter-drift",
                      "bundle_ref": bundle_ref, "source_ref": ref}, "resolve-source", expected=3)
    assert response["diagnostics"][0]["code"] == "SOURCE_EXTRACTION_UNAVAILABLE"


@pytest.mark.parametrize("kind", ["pdf", "docx"])
def test_ci_native_document_check_still_requires_explicit_original_input(tmp_path, kind):
    p = pipeline(tmp_path)
    data, snapshot, lock, ref = frozen_source(kind)
    source_path = "docs-raw/specs/native." + kind
    p.write(source_path, data)
    lock["schema_version"] = "2.0"
    lock["resources"][0]["path"] = source_path
    snapshot["recovery"] = {"method": "evidence_object", "path": "logs/evidence/objects/missing.bin"}
    p.write("governance/sources.lock.yaml", yaml.safe_dump(lock, sort_keys=False).encode())
    raw = p.card("native-rule", {"snapshots": [snapshot]})
    _, front, body = raw.decode().split("---", 2)
    metadata = yaml.safe_load(front)
    metadata["sources"][0]["locator"] = ref["locator"]
    p.write("knowledge/rules/native-rule.md", ("---\n" + yaml.safe_dump(metadata, sort_keys=False) + "---" + body).encode())
    p.write("knowledge/rules/index.md", b"# Rules\n[native](native-rule.md)\n")
    framework = p.root / "governance/framework.lock.yaml"
    rules = v.load_rules(p.root, v.sha256_bytes(framework.read_bytes()))
    view = checks.prepare_view(p.root, mode="ci")
    missing = checks.run(view, rules)
    source_check = next(row for row in missing["checks"] if row["group"] == "source")
    assert source_check["status"] == "incomplete"
    assert any(row["code"] == "SOURCE_INPUT_UNAVAILABLE" for row in missing["findings"])
    supplied = checks.run(view, rules, {"source_objects": {snapshot["snapshot_id"]: data}})
    source_check = next(row for row in supplied["checks"] if row["group"] == "source")
    assert source_check["status"] == "passed", supplied["findings"]


def test_git_document_capture_preserves_its_local_object_fallback(tmp_path):
    p = pipeline(tmp_path)
    workspace = p.root.parent
    config = {"schema_version": "2.0", "workspace_id": "native-fixture",
              "knowledge_roots": {"wifi": p.root.name}, "code_root": "main", "modules": []}
    (workspace / "config.yml").write_text(yaml.safe_dump(config), encoding="utf-8")
    (workspace / "main").mkdir()
    original = b"# Original\nCapture marker is fixed.\n"
    source_path = "docs-raw/specs/retained.md"
    p.write(source_path, original)
    FIXTURE.git(workspace, "init", "-b", "main")
    FIXTURE.git(workspace, "add", "config.yml", p.root.name + "/" + source_path)
    FIXTURE.git(workspace, "commit", "-m", "Isolated raw original")
    bundle_ref = p.call("knowledge-source-capture", {"operation_id": "git-fallback",
                       "source_path": source_path, "source_id": "document"})["data"]["bundle_ref"]
    snapshot = p.record(bundle_ref)["snapshots"][0]
    assert snapshot["recovery"]["kind"] == "git"
    assert snapshot["local_recovery"]["method"] == "evidence_object"
    assert (p.root / snapshot["local_recovery"]["path"]).read_bytes() == original
