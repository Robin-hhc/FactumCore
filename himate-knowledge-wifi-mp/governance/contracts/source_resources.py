"""Exact logical-resource/snapshot resolution; source objects are always read-only."""

from .validation import ContractError, canonical_bytes, load_yaml, safe_path, sha256_bytes, workspace_config_path


def extract_document(data, suffix, *, primary=True):
    """The fixed document adapter shared by acquisition and read-only checks."""
    import io
    if suffix in {".md", ".txt"}:
        text = data.decode("utf-8-sig")
        if "\0" in text or not text.strip():
            raise ContractError(
                "UNREADABLE_DOCUMENT", "UTF-8 text must be nonempty and contain no NUL"
            )
        lines = text.splitlines()
        return (
            text,
            [
                {"kind": "lines", "start": i + 1, "end": i + 1, "text": line}
                for i, line in enumerate(lines)
            ],
            "utf8",
            [],
        )
    if suffix == ".pdf":
        from pypdf import PdfReader, __version__

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ContractError("ENCRYPTED_DOCUMENT", "Encrypted PDF is unsupported")
        pages = [page.extract_text() or "" for page in reader.pages]
        if not pages or any(not page.strip() for page in pages):
            raise ContractError(
                "PARTIAL_DOCUMENT",
                "PDF page lacks a readable text layer; no OCR",
            )
        return (
            "\n\n".join(pages),
            [
                {"kind": "page", "page": i + 1, "text": text}
                for i, text in enumerate(pages)
            ],
            "pypdf:" + __version__,
            ["PDF image content and visual layout are not interpreted"],
        )
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = []
        for i, paragraph in enumerate(document.paragraphs):
            if paragraph.text.strip():
                parts.append(
                    {"kind": "paragraph", "paragraph": i + 1, "text": paragraph.text}
                )
        for i, table in enumerate(document.tables):
            for j, row in enumerate(table.rows):
                parts.append(
                    {
                        "kind": "table_row",
                        "table": i + 1,
                        "row": j + 1,
                        "text": " | ".join(cell.text for cell in row.cells),
                    }
                )
        if not parts or not any(p["text"].strip() for p in parts):
            raise ContractError(
                "PARTIAL_DOCUMENT",
                "DOCX has no readable paragraphs or tables",
            )
        if document.element.xpath(".//w:ins | .//w:del"):
            raise ContractError(
                "PARTIAL_DOCUMENT",
                "DOCX tracked changes require explicit resolution",
            )
        return (
            "\n\n".join(p["text"] for p in parts),
            parts,
            "python-docx:" + docx.__version__,
            [
                "Embedded images, floating text, headers, and visual table layout are not interpreted"
            ],
        )
    if primary:
        raise ContractError(
            "UNSUPPORTED_DOCUMENT_FORMAT", "Supported primary formats: MD TXT PDF DOCX"
        )
    return (
        None,
        [],
        "opaque",
        ["Attachment bytes retained; content was not interpreted"],
    )
def extract_code_text(data, extractor):
    """Decode an explicitly fixed code adapter without changing original bytes."""
    codecs = {"code-utf8:1": "utf-8-sig", "code-gb18030:1": "gb18030"}
    if extractor not in codecs:
        raise ContractError("SOURCE_ADAPTER_UNAVAILABLE", "Unknown fixed code extraction adapter")
    try:
        text = data.decode(codecs[extractor], errors="strict")
        if "\0" in text:
            raise UnicodeError("NUL in code text")
    except UnicodeError as exc:
        raise ContractError("SOURCE_ADAPTER_UNAVAILABLE", "Exact code bytes cannot be decoded by the declared adapter") from exc
    locators = [{"kind": "lines", "start": index + 1, "end": index + 1, "text": line}
                for index, line in enumerate(text.splitlines())]
    return text, locators


def read_source_lock(knowledge_root):
    lock = load_yaml(safe_path(knowledge_root, "governance/sources.lock.yaml"))
    if lock.get("schema_version") != "2.0":
        raise ContractError("UNSUPPORTED_SOURCE_VERSION", "Expected 2.0")
    for field, key in (("resources", "source_id"), ("snapshots", "snapshot_id")):
        identities = [item[key] for item in lock[field]]
        if len(identities) != len(set(identities)):
            raise ContractError("DUPLICATE_SOURCE", field)
    resource_ids = {item["source_id"] for item in lock["resources"]}
    if any(item["source_id"] not in resource_ids for item in lock["snapshots"]):
        raise ContractError(
            "SOURCE_RESOURCE_MISSING", "Snapshot points to an unregistered resource"
        )
    return lock


def validate_source_ref(ref, lock, roots=None):
    if not isinstance(ref, dict) or not all(
        ref.get(k) for k in ("evidence_ref", "source_id", "snapshot_id", "locator")
    ):
        raise ContractError(
            "INCOMPLETE_EVIDENCE_REF", "Evidence, source, snapshot and locator required"
        )
    matches = [s for s in lock["snapshots"] if s["snapshot_id"] == ref["snapshot_id"]]
    if len(matches) != 1 or matches[0]["source_id"] != ref["source_id"]:
        raise ContractError("SOURCE_SNAPSHOT_MISMATCH", ref["snapshot_id"])
    if not any(s["source_id"] == ref["source_id"] for s in lock["resources"]):
        raise ContractError("SOURCE_RESOURCE_MISSING", ref["source_id"])
    return matches[0]


def validate_snapshot(bundle, expected=None):
    for key, value in (expected or {}).items():
        if bundle.get(key) != value:
            raise ContractError("SNAPSHOT_BINDING_MISMATCH", key)
    seen = set()
    for snapshot in bundle.get("snapshots", []):
        identity = snapshot.get("snapshot_id")
        if not identity or identity in seen:
            raise ContractError("DUPLICATE_SNAPSHOT", identity)
        if (
            not snapshot.get("sha256")
            or not isinstance(snapshot.get("bytes"), int)
            or snapshot["bytes"] < 0
        ):
            raise ContractError("INCOMPLETE_SNAPSHOT", identity)
        if not snapshot.get("recovery"):
            raise ContractError("UNRECOVERABLE_SNAPSHOT", identity)
        seen.add(identity)
    return bundle


def recover_snapshot(knowledge_root, snapshot, *, allow_local_fallback=True):
    return recover_snapshot_source(knowledge_root, snapshot, allow_local_fallback=allow_local_fallback)[0]


def recover_snapshot_source(knowledge_root, snapshot, *, allow_local_fallback=True):
    """Recover fixed bytes and report the actual source, including clean checkouts."""
    recovery = snapshot.get("recovery", {})
    if isinstance(recovery, dict) and (recovery.get("kind") == "git" or recovery.get("method") == "git"):
        try:
            data = _recover_git(knowledge_root, recovery)
        except ContractError as exc:
            unavailable = {"RECOVERY_CONFIG_REQUIRED", "RECOVERY_CONFIG_MISMATCH", "RECOVERY_REPO_UNKNOWN", "RECOVERY_REPO_MISMATCH", "GIT_SOURCE_UNAVAILABLE", "INVALID_YAML"}
            if not allow_local_fallback or exc.code not in unavailable:
                raise
            local = snapshot.get("local_recovery")
            data = None
            if isinstance(local, dict) and local.get("method") == "evidence_object":
                try:
                    data = safe_path(knowledge_root, local["path"], ["logs/evidence"]).read_bytes()
                    reference = local["path"]
                except FileNotFoundError:
                    pass
            if data is None:
                current = _current_registered_source(knowledge_root, snapshot)
                if current is not None:
                    return current
                raise
        else:
            reference = "git:" + recovery.get("commit", recovery.get("revision", "")) + ":" + recovery["path"]
        if sha256_bytes(data) != snapshot.get("sha256") or len(data) != snapshot.get("bytes"):
            raise ContractError("EVIDENCE_OBJECT_CHANGED", snapshot.get("snapshot_id"))
        return data, reference
    path = recovery.get("path") if isinstance(recovery, dict) else recovery
    if not path:
        raise ContractError("UNSUPPORTED_RECOVERY", "Immutable local object required")
    try:
        data = safe_path(knowledge_root, path, ["logs/evidence"]).read_bytes()
    except FileNotFoundError:
        current = _current_registered_source(knowledge_root, snapshot) if allow_local_fallback else None
        if current is not None:
            return current
        raise ContractError("SOURCE_INPUT_UNAVAILABLE", "Exact snapshot is unavailable in retained objects, registered current files, or local Git history (at most 1000 path changes; shallow clones may omit the version)")
    if sha256_bytes(data) != snapshot.get("sha256") or len(data) != snapshot.get(
        "bytes"
    ):
        raise ContractError("EVIDENCE_OBJECT_CHANGED", snapshot.get("snapshot_id"))
    return data, path


def _current_registered_source(knowledge_root, snapshot):
    """A shared lock may recover a current file only when its exact bytes match.

    No scan, alternate path, or newer source substitutes for a missing version.
    Candidate bundles still use retained objects until their lock is installed.
    """
    from pathlib import Path
    from .validation import _workspace_no_links, normalize_workspace, workspace_file_allowed

    try:
        lock = read_source_lock(knowledge_root)
    except FileNotFoundError:
        return None
    registered = [row for row in lock["snapshots"] if row.get("snapshot_id") == snapshot.get("snapshot_id")]
    if len(registered) != 1 or any(registered[0].get(key) != snapshot.get(key) for key in ("source_id", "sha256", "bytes")):
        return None
    resources = [row for row in lock["resources"] if row.get("source_id") == snapshot.get("source_id")]
    if len(resources) != 1:
        return None
    resource = resources[0]
    if resource.get("kind") in {"raw", "attachment"}:
        allowed = ["docs-raw/" + name for name in ("conventions", "chips", "storys", "specs", "protocols", "troubleshooting")]
        path = safe_path(knowledge_root, resource["path"], allowed)
        reference = resource["path"]
    elif resource.get("kind") == "code":
        config_path = workspace_config_path(knowledge_root)
        workspace = config_path.parent
        configuration = normalize_workspace(load_yaml(config_path), workspace)
        rows = [row for row in configuration["repos"] if row["repo_id"] == resource.get("repo_id")]
        if len(rows) != 1:
            return None
        if not workspace_file_allowed(configuration, resource["repo_id"], resource["path"]):
            raise ContractError("RECOVERY_SCOPE_REQUIRED", "Recovered code must remain within its configured module scope")
        declared = Path(rows[0]["root"])
        repository = declared if declared.is_absolute() else safe_path(workspace, declared.as_posix())
        path = safe_path(repository, resource["path"])
        reference = "code:" + resource["repo_id"] + ":" + resource["path"]
    else:
        return None
    _workspace_no_links(path)
    try:
        data = path.read_bytes()
    except FileNotFoundError:
        return _registered_raw_history(knowledge_root, resource, snapshot) if resource.get("kind") in {"raw", "attachment"} else None
    if sha256_bytes(data) != snapshot.get("sha256") or len(data) != snapshot.get("bytes"):
        if resource.get("kind") in {"raw", "attachment"}:
            return _registered_raw_history(knowledge_root, resource, snapshot)
        return None
    return data, reference


def _registered_raw_history(knowledge_root, resource, snapshot):
    """Search only this registered raw path in bounded, already local history."""
    from pathlib import Path

    root = Path(knowledge_root).resolve()
    try:
        workspace = workspace_config_path(root).parent
        if not (workspace / ".git").exists():
            return None
        top = Path(_git_read(workspace, "rev-parse", "--show-toplevel").decode().strip()).resolve()
        if top != workspace:
            return None
        relative = root.relative_to(workspace).as_posix() + "/" + resource["path"]
        commits = _git_read(workspace, "log", "--no-show-signature", "--format=%H", "--max-count=1001", "HEAD", "--", relative).decode("ascii").splitlines()
        for commit in commits[:1000]:
            try:
                data = _git_read(workspace, "cat-file", "blob", commit + ":" + relative)
            except ContractError as exc:
                if exc.code == "GIT_SOURCE_UNAVAILABLE":
                    continue
                raise
            if len(data) == snapshot.get("bytes") and sha256_bytes(data) == snapshot.get("sha256"):
                return data, "git:" + commit + ":" + relative
    except (ContractError, OSError, UnicodeError):
        return None
    return None


def _recover_git(knowledge_root, recovery):
    """Read a pinned object only from the configured workspace/repository.

    A source cannot choose a repository path. Raw Git paths are relative to the
    co-located workspace; code Git paths need an explicit configured repo_id.
    """
    from pathlib import Path, PurePosixPath
    import re

    root = Path(knowledge_root).resolve()
    try:
        config_path = workspace_config_path(root)
    except ContractError as exc:
        raise ContractError("RECOVERY_CONFIG_REQUIRED", "Git recovery requires one exact workspace config") from exc
    workspace = config_path.parent
    config = load_yaml(config_path)
    if not isinstance(config, dict):
        raise ContractError("RECOVERY_CONFIG_MISMATCH", "Configured workspace mapping is unavailable")
    mapped = config.get("knowledge_roots", {}).get("wifi")
    if not isinstance(mapped, str) or (workspace / mapped).resolve() != root:
        raise ContractError("RECOVERY_CONFIG_MISMATCH", "Configured domain does not match this knowledge root")
    if any(key in recovery for key in ("repo_root", "repository_root", "workspace_root")):
        raise ContractError("UNTRUSTED_RECOVERY_ROOT", "Repository roots come only from config")
    commit = recovery.get("commit", recovery.get("revision"))
    relative = recovery.get("path")
    parts = PurePosixPath(relative).parts if isinstance(relative, str) else ()
    if not isinstance(commit, str) or not re.fullmatch(r"[0-9a-fA-F]{40}|[0-9a-fA-F]{64}", commit):
        raise ContractError("EXACT_GIT_REVISION_REQUIRED", "A full immutable commit object ID is required")
    if not parts or any(part in {".", "..", ".git"} for part in parts) or relative.startswith("/") or "\\" in relative or ":" in relative:
        raise ContractError("UNSAFE_RECOVERY_PATH", "Expected an exact repository-relative file")
    repo_id = recovery.get("repo_id")
    if repo_id:
        from .validation import normalize_workspace, workspace_file_allowed
        try:
            configured = normalize_workspace(config, workspace) if "code_root" in config or "modules" in config else config
        except (ContractError, OSError) as exc:
            raise ContractError("RECOVERY_CONFIG_MISMATCH", "Current module layout is unavailable for fixed Git recovery") from exc
        matches = [r for r in configured.get("repos", []) if r.get("repo_id") == repo_id and r.get("root")]
        if len(matches) != 1:
            raise ContractError("RECOVERY_REPO_UNKNOWN", str(repo_id))
        if not workspace_file_allowed(configured, repo_id, relative):
            raise ContractError("RECOVERY_SCOPE_REQUIRED", "Git code recovery must remain within its configured module scope")
        repository = Path(matches[0]["root"])
        repository = repository.resolve() if repository.is_absolute() else (workspace / repository).resolve()
    else:
        prefix = root.relative_to(workspace).as_posix() + "/docs-raw/"
        tail = relative[len(prefix):] if relative.startswith(prefix) else ""
        raw_classes = {"conventions", "chips", "storys", "specs", "protocols", "troubleshooting"}
        if tail.split("/", 1)[0] not in raw_classes or "/" not in tail:
            raise ContractError("RECOVERY_SCOPE_REQUIRED", "Default repository recovery requires a supported raw original")
        repository = workspace
    safe_path(repository, relative)
    def git(*args):
        return _git_read(repository, *args)

    top = Path(git("rev-parse", "--show-toplevel").decode().strip()).resolve()
    if top != repository:
        raise ContractError("RECOVERY_REPO_MISMATCH", "Configured repository is not the exact Git root")
    if git("cat-file", "-t", commit).strip() != b"commit":
        raise ContractError("EXACT_GIT_REVISION_REQUIRED", "Recovery object must be a commit")
    return git("cat-file", "blob", commit + ":" + relative)


def _git_read(repository, *args):
    import os
    import subprocess
    if not args or args[0] not in {"rev-parse", "cat-file", "log"}:
        raise ContractError("UNSAFE_GIT_READ", "Only fixed object/ref reads are supported")
    env = {key: value for key, value in os.environ.items() if not key.startswith("GIT_") or key == "GIT_TEST_ASSUME_DIFFERENT_OWNER"}
    env.update({"GIT_CONFIG_NOSYSTEM": "1", "GIT_CONFIG_GLOBAL": os.devnull, "GIT_TERMINAL_PROMPT": "0",
                "GIT_OPTIONAL_LOCKS": "0", "GIT_PAGER": "cat", "GIT_NO_LAZY_FETCH": "1"})
    command = ["git", "--no-optional-locks", "-c", "core.fsmonitor=false", "-c", "core.hooksPath=" + os.devnull,
               "-c", "diff.external=", "-c", "core.pager=cat", "-c", "log.showSignature=false", "-c", "protocol.allow=never", "-c", "protocol.file.allow=never",
               "-c", "safe.directory=" + str(repository.resolve()).replace("\\", "/"), "-C", str(repository), *args]
    try:
        process = subprocess.run(command, capture_output=True, env=env, check=False, timeout=30)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ContractError("GIT_SOURCE_UNAVAILABLE", "Exact configured Git source read unavailable") from exc
    if process.returncode:
        raise ContractError("GIT_SOURCE_UNAVAILABLE", "Exact configured Git source object is unavailable")
    return process.stdout


def git_recovery_for_current(knowledge_root, source_path, expected_bytes):
    """Discover a portable raw version only when configured HEAD has exact bytes."""
    from pathlib import Path

    root = Path(knowledge_root).resolve()
    try:
        config_path = workspace_config_path(root)
    except (OSError, ContractError):
        return None
    workspace = config_path.parent
    if not (workspace / ".git").exists():
        return None
    try:
        commit = _git_read(workspace, "rev-parse", "--verify", "HEAD").decode().strip()
    except ContractError:
        return None
    recovery = {"method": "git", "kind": "git", "commit": commit,
                "path": root.relative_to(workspace).as_posix() + "/" + source_path, "status": "available"}
    try:
        return recovery if _recover_git(root, recovery) == expected_bytes else None
    except (OSError, ContractError):
        return None


def document_segments(addresses, locator, suffix):
    """Select native segments identically for fixed checks and source reads."""
    kind = locator["kind"]
    if suffix == ".pdf" and kind in {"page", "pages"}:
        address_kind, field = "page", "page"
    elif suffix == ".docx" and kind in {"paragraph", "paragraphs"}:
        address_kind, field = "paragraph", "paragraph"
    elif suffix == ".docx" and kind == "table_row":
        table, row = locator.get("table"), locator.get("row")
        if any(isinstance(value, bool) or not isinstance(value, int) or value < 1 for value in (table, row)):
            raise ContractError("INVALID_SOURCE_LOCATOR", "Native table and row addresses must be positive integers")
        selected = [item for item in addresses if item["kind"] == "table_row"
                    and item["table"] == table and item["row"] == row]
        if len(selected) != 1:
            raise ContractError("INVALID_SOURCE_LOCATOR", "Native table row is outside the exact original")
        return selected
    else:
        raise ContractError("INVALID_SOURCE_LOCATOR", "This document format requires a native page, paragraph or table-row address")
    first = locator.get(field, locator.get("start", locator.get("number", locator.get("id"))))
    last = locator.get("end", first)
    if (isinstance(first, bool) or isinstance(last, bool) or not isinstance(first, int)
            or not isinstance(last, int) or not 1 <= first <= last):
        raise ContractError("INVALID_SOURCE_LOCATOR", "Native document address must be a positive integer range")
    available = {row[field] for row in addresses if row["kind"] == address_kind}
    if not available or last > max(available) or any(value not in available for value in range(first, last + 1)):
        raise ContractError("INVALID_SOURCE_LOCATOR", "Native document address is outside the exact original")
    return [row for row in addresses if row["kind"] == address_kind and first <= row[field] <= last]


def _validate_document_reference(snapshot, data, locator):
    """Replay frozen native document addresses from the already verified bytes."""
    suffix, retained = snapshot.get("format"), snapshot.get("extraction")
    if suffix not in {".pdf", ".docx"} or not retained:
        return False
    try:
        text, addresses, extractor, _ = extract_document(data, suffix)
    except (ImportError, ContractError) as exc:
        raise ContractError("SOURCE_ADAPTER_UNAVAILABLE", "The fixed document adapter cannot recover native addresses") from exc
    if extractor != retained.get("extractor"):
        raise ContractError("SOURCE_ADAPTER_UNAVAILABLE", "The installed document adapter differs from the frozen adapter")
    frozen_locator = sha256_bytes(canonical_bytes({"locators": addresses}))
    if (sha256_bytes(text.encode("utf-8")) != retained.get("text_sha256")
            or frozen_locator != retained.get("locator_sha256")
            or snapshot.get("locator", {}).get("sha256", frozen_locator) != frozen_locator):
        raise ContractError("SOURCE_EXTRACTION_UNAVAILABLE", "The exact original cannot reproduce frozen document text and addresses")
    document_segments(addresses, locator, suffix)
    return True


def validate_reference(ref, lock, resolver, *, allowed_kinds=None):
    """Recover an exact object and validate its address, without semantic inference.

    Resolvers return bytes, or {bytes, locators}; structured locators must be
    derived from those exact bytes by the caller's fixed extraction adapter.
    Unavailable adapters raise an explicit incomplete error, never a pass.
    """
    snapshot = validate_source_ref(ref, lock)
    locator = ref["locator"]
    kinds = allowed_kinds or {"lines", "page", "pages", "paragraph", "paragraphs", "table", "table_row", "attachment", "section"}
    if not isinstance(locator, dict) or locator.get("kind") not in kinds:
        raise ContractError("INVALID_SOURCE_LOCATOR", "Unknown locator kind")
    try:
        resolved = resolver(snapshot)
    except (OSError, KeyError) as exc:
        raise ContractError("SOURCE_INPUT_UNAVAILABLE", "Required exact source is unavailable") from exc
    data = resolved.get("bytes") if isinstance(resolved, dict) else resolved
    if not isinstance(data, bytes):
        raise ContractError("SOURCE_INPUT_UNAVAILABLE", "Required exact source bytes are unavailable")
    if sha256_bytes(data) != snapshot.get("sha256") or len(data) != snapshot.get("bytes"):
        raise ContractError("EVIDENCE_OBJECT_CHANGED", "Exact source fingerprint differs")
    kind = locator["kind"]
    code = any(item.get("source_id") == snapshot["source_id"] and item.get("kind") == "code" for item in lock["resources"])
    if not code and _validate_document_reference(snapshot, data, locator):
        pass
    elif kind in {"lines", "paragraph", "paragraphs", "section"}:
        extraction = snapshot.get("extraction")
        if code and extraction:
            text, addresses = extract_code_text(data, extraction.get("extractor"))
            if (sha256_bytes(text.encode("utf-8")) != extraction.get("text_sha256")
                    or sha256_bytes(canonical_bytes({"locators": addresses})) != extraction.get("locator_sha256")):
                raise ContractError("SOURCE_EXTRACTION_CHANGED", "Fixed code text or locator fingerprint differs")
        else:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeError as exc:
                raise ContractError("SOURCE_ADAPTER_UNAVAILABLE", "Text locator requires a fixed extraction adapter") from exc
        if kind == "section":
            from .validation import markdown_structure
            anchor = locator.get("anchor") or locator.get("id")
            if not anchor or str(anchor).lstrip("#") not in markdown_structure(text)["anchors"]:
                raise ContractError("INVALID_SOURCE_LOCATOR", "Source section locator does not exist")
        else:
            import re
            segments = text.splitlines() if kind == "lines" else re.split(r"\n\s*\n", text.strip())
            first = locator.get("start", locator.get("number"))
            last = locator.get("end", first)
            if isinstance(first, bool) or isinstance(last, bool) or not isinstance(first, int) or not isinstance(last, int) or not 1 <= first <= last <= len(segments):
                raise ContractError("INVALID_SOURCE_LOCATOR", "Source locator is outside the exact object")
    else:
        # A caller cannot substitute an unbound page/table list for the object.
        # The adapter must declare the same byte hash, plus readable addresses.
        if not isinstance(resolved, dict) or resolved.get("extracted_from_sha256") != snapshot["sha256"]:
            raise ContractError("SOURCE_ADAPTER_UNAVAILABLE", "Structured locator requires exact-byte extraction")
        addresses = resolved.get("locators", {}).get(kind)
        identity = locator.get("id", locator.get("number", locator.get("start")))
        end = locator.get("end", identity)
        wanted = list(range(identity, end + 1)) if isinstance(identity, int) and isinstance(end, int) else [identity]
        if not isinstance(addresses, (list, dict)) or not wanted or any(value not in addresses for value in wanted):
            raise ContractError("INVALID_SOURCE_LOCATOR", "Structured source locator is not readable")
    return {"snapshot_id": snapshot["snapshot_id"], "sha256": snapshot["sha256"], "locator": locator}


def validate_code_references(card, lock, checkout, rules):
    """Check fixed code addresses against job-supplied Git, never support claims.

    checkout = {configuration: exact config.yml path, config_digest: SHA,
                repos: {repo_id: {root: absolute checkout, commit: fixed SHA}}}.
    The configuration is the sole mapping authority. No path comes from raw prose.
    """
    import os
    from pathlib import Path
    import re
    import subprocess
    from .validation import normalize_workspace, validate_shape, workspace_file_allowed
    if not isinstance(checkout, dict) or not checkout.get("configuration") or not checkout.get("repos"):
        raise ContractError("CODE_INPUT_UNAVAILABLE", "Fixed checkout and exact configuration are required")
    configuration = Path(checkout["configuration"])
    if not configuration.is_absolute() or configuration.is_symlink():
        raise ContractError("CODE_INPUT_UNAVAILABLE", "Absolute unlinked configuration is required")
    data = configuration.read_bytes()
    if sha256_bytes(data) != checkout.get("config_digest"):
        raise ContractError("CODE_CONFIGURATION_MISMATCH", "Fixed configuration fingerprint differs")
    config = load_yaml(data)
    validate_shape("workspace", config, rules)
    config = normalize_workspace(config, configuration.parent)
    registered = {item["repo_id"]: item for item in config["repos"]}
    targets = {item["target_id"]: item for item in config["targets"]}
    for context in card["metadata"]["scope"]["contexts"]:
        target_id = context.get("target_id")
        if target_id:
            target = targets.get(target_id)
            fields = ("chip", "side") if config.get("mapping_kind") == "modules" else ("chip", "side", "build_profile", "pairing_id")
            if not target or any(context.get(key) and context[key] != target.get(key) for key in fields):
                raise ContractError("CODE_TARGET_MISMATCH", "Target/build reference differs from configuration")
    resources = {item["source_id"]: item for item in lock["resources"]}
    refs = [ref for ref in card["metadata"]["sources"] if resources.get(ref["source_id"], {}).get("kind") == "code"]
    if not refs:
        raise ContractError("CODE_INPUT_UNAVAILABLE", "A code card requires explicit fixed code references")
    environment = {key: value for key, value in os.environ.items() if not key.startswith("GIT_") or key == "GIT_TEST_ASSUME_DIFFERENT_OWNER"}
    environment.update(GIT_CONFIG_NOSYSTEM="1", GIT_CONFIG_GLOBAL=os.devnull, GIT_TERMINAL_PROMPT="0", GIT_OPTIONAL_LOCKS="0", GIT_NO_LAZY_FETCH="1")
    verified = []
    for ref in refs:
        resource, locator = resources[ref["source_id"]], ref["locator"]
        repo_id = locator.get("repo_id") or resource.get("repo_id")
        relative = locator.get("path") or resource.get("path")
        job_repo, mapping = checkout["repos"].get(repo_id), registered.get(repo_id)
        if not job_repo or not mapping or not mapping.get("root") or mapping["root"] == "unknown":
            raise ContractError("CODE_INPUT_UNAVAILABLE", "Code repository mapping is unavailable")
        root = Path(job_repo.get("root", ""))
        declared = Path(mapping["root"])
        if not declared.is_absolute():
            declared = configuration.parent / declared
        if not root.is_absolute() or root.resolve() != declared.resolve():
            raise ContractError("CODE_REPOSITORY_MISMATCH", "Checkout differs from the sole configured repository mapping")
        if config.get("mapping_kind") == "modules":
            target_ids = [context.get("target_id") for context in card["metadata"]["scope"]["contexts"]] or [None]
            if not any(workspace_file_allowed(config, repo_id, relative, target_id) for target_id in target_ids):
                raise ContractError("CODE_PATH_MISMATCH", "Code reference is outside the configured module/Target range", relative)
            raw = safe_path(root, relative).read_bytes()
            snapshot = validate_source_ref(ref, lock)
            if sha256_bytes(raw) != snapshot["sha256"] or len(raw) != snapshot["bytes"]:
                raise ContractError("CODE_FINGERPRINT_MISMATCH", "Current fixed code differs; return to local ingest")
            if job_repo.get("files") and job_repo["files"].get(relative) != sha256_bytes(raw):
                raise ContractError("CODE_FINGERPRINT_MISMATCH", "Job file fingerprint differs from actual source")
            validate_reference(ref, lock, lambda snapshot: raw)
            verified.append({"repo_id": repo_id, "path": relative, "commit": None,
                             "sha256": snapshot["sha256"], "evidence_basis": "filesystem"})
            continue
        safe_path(root, relative)
        commit = job_repo.get("commit")
        if not re.fullmatch(r"[0-9a-f]{40}|[0-9a-f]{64}", commit or ""):
            raise ContractError("CODE_INPUT_UNAVAILABLE", "An exact code commit is required")
        def git(arguments):
            command = ["git", "--no-optional-locks", "-c", "core.fsmonitor=false", "-c", "core.hooksPath=" + os.devnull, "-c", "core.pager=cat", "-c", "protocol.allow=never", "-c", "safe.directory=" + str(root.resolve()).replace("\\", "/"), "-C", str(root), *arguments]
            result = subprocess.run(command, capture_output=True, timeout=30, env=environment)
            if result.returncode:
                raise ContractError("CODE_INPUT_UNAVAILABLE", "Fixed Git code object is unavailable")
            return result.stdout
        if git(["rev-parse", "HEAD"]).decode("ascii").strip() != commit:
            raise ContractError("CODE_COMMIT_MISMATCH", "Job checkout is not at its declared fixed commit")
        tree = git(["ls-tree", "-z", commit, "--", relative])
        if not tree.startswith((b"100644 ", b"100755 ")) or tree.count(b"\x00") != 1:
            raise ContractError("CODE_PATH_MISMATCH", "Code reference must be a regular tracked file")
        object_name = commit + ":" + relative
        if git(["cat-file", "-t", object_name]).strip() != b"blob":
            raise ContractError("CODE_PATH_MISMATCH", "Code reference is not a regular Git blob")
        raw = git(["cat-file", "blob", object_name])
        snapshot = validate_source_ref(ref, lock)
        if sha256_bytes(raw) != snapshot["sha256"] or len(raw) != snapshot["bytes"]:
            raise ContractError("CODE_FINGERPRINT_MISMATCH", "Current fixed code differs; return to local ingest")
        validate_reference(ref, lock, lambda snapshot: raw)
        verified.append({"repo_id": repo_id, "path": relative, "commit": commit, "sha256": snapshot["sha256"]})
    return verified


def _fixed_code_git(root, arguments):
    import os
    import subprocess
    allowed = arguments and arguments[0] in {"rev-parse", "cat-file", "ls-tree"}
    allowed = allowed or arguments == ["config", "--local", "--get", "remote.origin.url"]
    if not allowed:
        raise ContractError("UNSAFE_CODE_OPERATION", "Only fixed Git object/identity reads are allowed")
    environment = {key: value for key, value in os.environ.items() if not key.startswith("GIT_") or key == "GIT_TEST_ASSUME_DIFFERENT_OWNER"}
    environment.update(GIT_CONFIG_NOSYSTEM="1", GIT_CONFIG_GLOBAL=os.devnull, GIT_TERMINAL_PROMPT="0", GIT_OPTIONAL_LOCKS="0", GIT_NO_LAZY_FETCH="1", GIT_PAGER="cat")
    command = ["git", "--no-optional-locks", "-c", "core.fsmonitor=false", "-c", "core.hooksPath=" + os.devnull, "-c", "core.pager=cat", "-c", "protocol.allow=never", "-c", "protocol.file.allow=never", "-c", "safe.directory=" + str(root.resolve()).replace("\\", "/"), "-C", str(root), *arguments]
    try:
        result = subprocess.run(command, capture_output=True, timeout=30, env=environment)
    except (OSError, subprocess.SubprocessError) as exc:
        raise ContractError("CODE_INPUT_UNAVAILABLE", "Fixed code reader is unavailable") from exc
    if result.returncode:
        raise ContractError("CODE_INPUT_UNAVAILABLE", "Fixed Git code object or identity is unavailable")
    return result.stdout


def _validate_mapped_current_identity(identity, configuration, config, rules, read):
    """Recheck automatic repository facts using the same relative module map."""
    from pathlib import Path
    from .validation import workspace_file_allowed

    domain = Path(rules["knowledge_root"])
    try:
        workspace_config_path(domain, config_path=configuration)
    except ContractError as exc:
        raise ContractError("CODE_CONFIGURATION_CHANGED", "Module configuration does not map to the selected knowledge root") from exc
    source_basis = identity.get("basis") == "source"
    target, build_files = identity.get("target"), identity["build_files"]
    declared = None
    if target is not None:
        matches = [item for item in config["targets"] if item["target_id"] == target.get("target_id")]
        if len(matches) != 1:
            raise ContractError("CODE_TARGET_CHANGED", "Accepted Target is outside the configured modules")
        declared = matches[0]
        if any(target.get(field) != declared[field] for field in ("target_id", "chip", "side")):
            raise ContractError("CODE_TARGET_CHANGED", "Accepted Target differs from configured module labels")
    if not source_basis and (
        declared is None or not isinstance(target.get("features"), list)
        or not isinstance(target.get("protocol_versions"), list)
    ):
        raise ContractError("CODE_TARGET_UNAVAILABLE", "Actual Target context is incomplete")
    for relative, expected in build_files.items():
        if sha256_bytes(read(safe_path(domain, relative, ["logs/evidence"]))) != expected:
            raise ContractError("CODE_BUILD_CHANGED", "Accepted build artifact bytes changed", relative)
    registered = {item["repo_id"]: item for item in config["repos"]}
    for repo_id, item in identity["repos"].items():
        mapping = registered.get(repo_id)
        if not mapping or item.get("kind") != mapping["kind"] or item.get("workspace_path") != mapping["workspace_path"]:
            raise ContractError("CODE_REPOSITORY_CHANGED", "Accepted repository differs from the relative module mapping")
        root = (configuration.parent / mapping["root"]).resolve()
        kind = mapping["kind"]
        files = item.get("files")
        if not isinstance(files, dict) or not files:
            raise ContractError("CODE_IDENTITY_UNAVAILABLE", "Accepted relevant file manifest is missing")
        if kind != "export" or mapping.get("evidence_basis") != "filesystem" or item.get("evidence_basis") != "filesystem":
            raise ContractError("CODE_IDENTITY_UNAVAILABLE", "Module evidence requires its exact filesystem scope")
        if item.get("head") is not None or item.get("checked_files_dirty") is not None or item.get("identity") is not None:
            raise ContractError("CODE_IDENTITY_UNAVAILABLE", "Filesystem evidence cannot establish Git or clean-state facts")
        for relative, expected in files.items():
            if not workspace_file_allowed(config, repo_id, relative, target.get("target_id") if target else None):
                raise ContractError("CODE_TARGET_CHANGED", "Accepted file is outside the configured module/Target range", relative)
            current = read(safe_path(root, relative))
            if sha256_bytes(current) != expected:
                raise ContractError("CODE_BYTES_CHANGED", "Current relevant code bytes differ", relative)
    return identity


def validate_current_code_identity(identity, config_path, config_digest, rules, policy):
    """Recheck an accepted code identity against current protected inputs.

    This checks every accepted relevant file, including uncommitted bytes, the
    provided Target/build artifacts and exact Git identity. It never creates a new
    semantic proof or treats an acceptance summary as authentication.
    """
    import os
    from pathlib import Path
    import re
    from .validation import normalize_workspace, validate_shape

    def known(value):
        return value is not None and value not in ("", "unknown", "unavailable")

    def read(path):
        text = str(Path(path).absolute())
        if os.name == "nt" and not text.startswith("\\\\?\\"):
            text = "\\\\?\\UNC\\" + text[2:] if text.startswith("\\\\") else "\\\\?\\" + text
        try:
            return Path(text).read_bytes()
        except OSError as exc:
            raise ContractError("CODE_INPUT_UNAVAILABLE", "An accepted code/build file is unavailable") from exc

    if not isinstance(identity, dict) or identity.get("basis") not in (None, "source") or not identity.get("repos"):
        raise ContractError("CODE_IDENTITY_UNAVAILABLE", "Complete accepted code identity and build context are required")
    source_basis = identity.get("basis") == "source"
    target, build_files = identity.get("target"), identity.get("build_files")
    if (
        (target is not None and not isinstance(target, dict))
        or (not source_basis and not isinstance(target, dict))
        or not isinstance(build_files, dict)
        or (not source_basis and not build_files)
        or (build_files and not re.fullmatch(r"[0-9a-f]{64}", str(identity.get("macro_digest", ""))))
        or (not build_files and identity.get("macro_digest") is not None)
    ):
        raise ContractError("CODE_IDENTITY_UNAVAILABLE", "Accepted source/build identity is inconsistent")
    configuration = Path(config_path)
    if not configuration.is_absolute() or configuration.is_symlink():
        raise ContractError("CODE_INPUT_UNAVAILABLE", "An explicit actual configuration is required")
    raw = read(configuration)
    if sha256_bytes(raw) != config_digest:
        raise ContractError("CODE_CONFIGURATION_CHANGED", "Current configuration differs from accepted context")
    config = load_yaml(raw)
    validate_shape("workspace", config, rules)
    config = normalize_workspace(config, configuration.parent)
    if config.get("mapping_kind") == "modules":
        return _validate_mapped_current_identity(identity, configuration, config, rules, read)
    if not isinstance(policy, dict) or not policy or config.get("code_policy") != policy:
        raise ContractError("CODE_POLICY_UNAVAILABLE", "Protected code policy must equal the configuration declaration")
    declared = None
    if target is not None:
        matches = [item for item in config["targets"] if item["target_id"] == target.get("target_id")]
        if len(matches) != 1:
            raise ContractError("CODE_TARGET_CHANGED", "Accepted Target is not uniquely registered")
        declared = matches[0]
        for field in ("target_id", "chip", "side", "build_profile"):
            if (not source_basis and not known(target.get(field))) or (known(target.get(field)) and target[field] != declared.get(field)):
                raise ContractError("CODE_TARGET_CHANGED", "Current Target/build profile differs")
        if not source_basis and (not declared.get("module_roots") or not declared.get("source_bindings") or not known(target.get("revision")) or not isinstance(target.get("features"), list) or not isinstance(target.get("protocol_versions"), list)):
            raise ContractError("CODE_TARGET_UNAVAILABLE", "Target code bindings/features/protocol/revision are incomplete")
        if build_files and (not set(declared["required_features"]).issubset(target.get("features") or []) or not set(declared["protocol_versions"]).issubset(target.get("protocol_versions") or [])):
            raise ContractError("CODE_TARGET_CHANGED", "Required features or protocol versions differ")
    domain = Path(rules["knowledge_root"])
    for relative, expected in build_files.items():
        path = safe_path(domain, relative, ["logs/evidence"])
        if sha256_bytes(read(path)) != expected:
            raise ContractError("CODE_BUILD_CHANGED", "Accepted build artifact bytes changed", relative)
    registered = {item["repo_id"]: item for item in config["repos"]}
    for repo_id, item in identity["repos"].items():
        mapping = registered.get(repo_id)
        protected = policy.get("repo_roots", {}).get(repo_id)
        if not mapping or mapping.get("kind") != "git" or not known(mapping.get("root")) or not known(mapping.get("expected_identity")) or not protected:
            raise ContractError("CODE_REPOSITORY_UNAVAILABLE", "Current code root/identity is not authorized")
        root = Path(mapping["root"])
        if not root.is_absolute():
            root = configuration.parent / root
        root = root.resolve()
        if root != Path(protected).resolve() or root != Path(item.get("root", "")).resolve():
            raise ContractError("CODE_REPOSITORY_CHANGED", "Accepted repository path differs from protected mapping")
        actual_root = Path(_fixed_code_git(root, ["rev-parse", "--show-toplevel"]).decode("utf-8").strip()).resolve()
        origin = _fixed_code_git(root, ["config", "--local", "--get", "remote.origin.url"]).decode("utf-8").strip()
        head = _fixed_code_git(root, ["rev-parse", "--verify", "HEAD^{commit}"]).decode("ascii").strip()
        if actual_root != root or origin != mapping["expected_identity"] or head != item.get("head") or (target and known(target.get("revision")) and target["revision"] != head):
            raise ContractError("CODE_REPOSITORY_CHANGED", "Actual Git root/origin/HEAD differs from accepted identity")
        if not isinstance(item.get("files"), dict) or not item["files"] or not isinstance(item.get("checked_files_dirty"), bool):
            raise ContractError("CODE_IDENTITY_UNAVAILABLE", "Accepted relevant file manifest and dirty state are required")
        dirty = False
        for relative, expected in item["files"].items():
            if declared and declared["module_roots"] and not any(module == "." or relative == module or relative.startswith(module.rstrip("/") + "/") for module in declared["module_roots"]):
                raise ContractError("CODE_TARGET_CHANGED", "Accepted file is outside configured Target modules", relative)
            path = safe_path(root, relative)
            tree = _fixed_code_git(root, ["ls-tree", "-z", head, "--", relative])
            if not tree.startswith((b"100644 ", b"100755 ")) or tree.count(b"\x00") != 1:
                raise ContractError("CODE_PATH_CHANGED", "Accepted code must remain a regular tracked file", relative)
            current = read(path)
            if sha256_bytes(current) != expected:
                raise ContractError("CODE_BYTES_CHANGED", "Current relevant code bytes differ, including uncommitted changes", relative)
            committed = _fixed_code_git(root, ["cat-file", "blob", head + ":" + relative])
            dirty = dirty or current != committed
        if dirty != item["checked_files_dirty"]:
            raise ContractError("CODE_DIRTY_CHANGED", "Checked-file dirty state differs from accepted identity")
    return identity
