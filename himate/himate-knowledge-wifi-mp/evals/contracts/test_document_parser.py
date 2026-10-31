"""Fixed per-operation material bundles and real modern-format adapters."""
from io import BytesIO
from pathlib import Path
import socket
import sys
import uuid

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from support.contracts import module

parser = module("document_parser")
ContractError = module("validation").ContractError


def png_bytes():
    from PIL import Image
    stream = BytesIO()
    Image.new("RGB", (80, 60), "blue").save(stream, format="PNG")
    return stream.getvalue()


def two_frame_tiff_bytes():
    from PIL import Image
    stream = BytesIO()
    first = Image.new("RGB", (80, 60), "blue")
    second = Image.new("RGB", (80, 60), "red")
    first.save(stream, format="TIFF", save_all=True, append_images=[second])
    return stream.getvalue()


def vector_pdf_bytes(pages=2, size=None):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject, RectangleObject
    original = PdfReader(BytesIO(document_bytes(".pdf")))
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_page(original.pages[index % 2])
        if size is not None:
            page.mediabox = RectangleObject([0, 0, *size])
        content = DecodedStreamObject()
        content.set_data(page.get_contents().get_data() + b"\n0 0 1 RG 3 w 20 20 m 70 130 180 30 260 140 c S\n")
        page[NameObject("/Contents")] = writer._add_object(content)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


def document_bytes(suffix, *, pictures=False):
    stream = BytesIO()
    if suffix in {".md", ".txt"}:
        return b"# Requirement\n\nRetry three times.\n"
    if suffix == ".docx":
        from docx import Document
        document = Document()
        document.add_heading("Requirement", 1)
        document.add_paragraph("Retry three times.")
        if pictures:
            document.add_picture(BytesIO(png_bytes()))
        document.save(stream)
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        document = Workbook()
        sheet = document.active
        sheet.title = "Config"
        sheet.append(["Parameter", "Value"])
        sheet.append(["Retry count", "three"])
        if pictures:
            from openpyxl.drawing.image import Image
            sheet.add_image(Image(BytesIO(png_bytes())), "D4")
        document.save(stream)
    elif suffix == ".pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = "Requirement"
        slide.placeholders[1].text = "Retry three times."
        if pictures:
            slide.shapes.add_picture(BytesIO(png_bytes()), 100, 100)
        document.save(stream)
    elif suffix == ".pdf":
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        writer = PdfWriter()
        font = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
        for text in (b"Retry three times.", b"Second page requirement."):
            page = writer.add_blank_page(width=300, height=180)
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            content = DecodedStreamObject()
            content.set_data(b"BT /F1 12 Tf 20 140 Td (" + text + b") Tj ET")
            page[NameObject("/Contents")] = writer._add_object(content)
        writer.write(stream)
    elif suffix == ".html":
        return b"<html><body><h1>Requirement</h1><p>Retry three times.</p></body></html>"
    elif suffix == ".csv":
        return b"Parameter,Value\nRetry count,three\n"
    else:
        return png_bytes()
    return stream.getvalue()


def parse_bytes(root, suffix=".md", *, output=".build/materials/parsed", data=None, options=None, page_range=None):
    data = document_bytes(suffix) if data is None else data
    source = {"id": "synthetic-requirement", "path": "docs-raw/requirement" + suffix}
    descriptor = parser.parse(root, source, data, output, options=options, page_range=page_range)
    loaded = parser.load(root, descriptor, expected_source_id=source["id"], expected_sha256=parser.sha256_bytes(data))
    return descriptor, loaded


@pytest.mark.parametrize("suffix", [".md", ".txt", ".pdf", ".docx", ".xlsx", ".pptx", ".html", ".csv"])
def test_real_modern_documents_keep_readable_text_and_verified_materials(tmp_path, suffix):
    descriptor, loaded = parse_bytes(tmp_path, suffix)
    assert descriptor["parse_status"] == "parsed", descriptor["gaps"]
    assert "three" in loaded["text"]
    assert "three" in loaded["markdown"]
    assert descriptor["statistics"]["output_bytes"] > 0
    assert (tmp_path / descriptor["document_ref"]).is_file()
    assert loaded["manifest"]["profile"]["remote_services"] is False
    assert loaded["manifest"]["profile"]["generate_page_images"] is False


@pytest.mark.parametrize("suffix", [".docx", ".xlsx", ".pptx", ".png"])
def test_real_embedded_images_are_saved_and_bound_to_source(tmp_path, suffix):
    descriptor, loaded = parse_bytes(tmp_path, suffix, data=document_bytes(suffix, pictures=True))
    assert descriptor["images"], descriptor["gaps"]
    for image in descriptor["images"]:
        assert image["ref"]
        assert (tmp_path / image["ref"]).is_file()
        assert image["description_status"] == "agent_required"
        assert image["locator"]["source_sha256"] == descriptor["source_sha256"]
    assert "data:image" not in (tmp_path / descriptor["document_ref"]).read_text()
    assert "data:image" not in loaded["markdown"]
    if suffix == ".png":
        assert descriptor["parse_status"] == "partial"
        assert any(g["code"] == "OCR_REQUIRED" for g in descriptor["gaps"])


def test_two_frame_tiff_explicitly_reports_the_unprocessed_frame(tmp_path):
    descriptor, loaded = parse_bytes(tmp_path, ".tiff", data=two_frame_tiff_bytes())
    gap = next(g for g in descriptor["gaps"] if g["code"] == "UNSUPPORTED_MULTIFRAME")
    assert descriptor["parse_status"] == "partial"
    assert (gap["total_frames"], gap["read_frames"], gap["unread_frames"]) == (2, 1, 1)
    assert len(descriptor["images"]) == 1
    from PIL import Image
    with Image.open(tmp_path / descriptor["images"][0]["ref"]) as first:
        assert first.getpixel((0, 0)) == (0, 0, 255)
    with pytest.raises(ContractError, match="INVALID_SOURCE_LOCATOR"):
        parser.select(loaded, {"kind": "page", "page": 2})


def test_successful_standard_first_frame_does_not_hide_later_tiff_frames(tmp_path, monkeypatch):
    # Real TIFF decoding/export is exercised. Only the adapter's successful
    # status is simulated; this is a coverage gate, not an OCR quality claim.
    original = parser._docling
    def successful_first_frame(root, source, data, output, options, page_range, source_hash, profile_hash):
        text, segments, images, _, pages = original(root, source, data, output,
            {**options, "pipeline": "native"}, page_range, source_hash, profile_hash)
        return text, segments, images, [], pages
    monkeypatch.setattr(parser, "_docling", successful_first_frame)
    descriptor, _ = parse_bytes(tmp_path, ".tiff", data=two_frame_tiff_bytes(), options={"pipeline": "standard"})
    assert descriptor["parse_status"] == "partial"
    gap = next(g for g in descriptor["gaps"] if g["code"] == "UNSUPPORTED_MULTIFRAME")
    assert (gap["total_frames"], gap["read_frames"], gap["unread_frames"]) == (2, 1, 1)


def test_standard_missing_models_still_reports_all_unread_tiff_frames(tmp_path):
    descriptor, _ = parse_bytes(tmp_path, ".tiff", data=two_frame_tiff_bytes(), options={"pipeline": "standard"})
    gap = next(g for g in descriptor["gaps"] if g["code"] == "UNSUPPORTED_MULTIFRAME")
    assert (gap["total_frames"], gap["read_frames"], gap["unread_frames"]) == (2, 0, 2)
    assert any(g["code"] == "PARSER_MODELS_REQUIRED" for g in descriptor["gaps"])


def test_page_range_and_element_locations_use_original_identity(tmp_path):
    descriptor, loaded = parse_bytes(tmp_path, ".pdf", page_range=(2, 2))
    assert loaded["manifest"]["pages"] == [2]
    assert "Second page" in parser.select(loaded, {"kind": "page", "page": 2})[0]["text"]
    with pytest.raises(ContractError, match="INVALID_SOURCE_LOCATOR"):
        parser.select(loaded, {"kind": "page", "page": 1})
    element = next(row["locator"] for row in loaded["locators"]["segments"] if row["locator"]["kind"] == "element")
    assert parser.select(loaded, element)
    assert parser.select(loaded, {**element, "parser": {"threads": 4}}) == parser.select(loaded, element)
    with pytest.raises(ContractError, match="INVALID_SOURCE_LOCATOR"):
        parser.select(loaded, {**element, "profile_sha256": "0" * 64})
    assert descriptor["source_sha256"] == element["source_sha256"]


def test_scoped_vector_pdf_has_a_bounded_physical_page_preview(tmp_path):
    descriptor, loaded = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes(), page_range=(2, 2))
    previews = [image for image in descriptor["images"] if image["role"] == "page_preview"]
    assert len(previews) == 1
    assert previews[0]["locator"] == {"kind": "page", "page": 2}
    assert previews[0]["element_ref"] is None
    from PIL import Image
    with Image.open(tmp_path / previews[0]["ref"]) as image:
        assert image.size == (450, 270)
        pixels = image.convert("RGB").tobytes()
        assert any(pixels[n] < 50 and pixels[n + 1] < 50 and pixels[n + 2] > 200
                   for n in range(0, len(pixels), 3)), "The source's blue vector curve is missing"
    preview = loaded["manifest"]["page_preview"]
    assert preview["rendered_pages"] == [2]
    assert preview["source_sha256"] == descriptor["source_sha256"]
    assert "images/page_preview_000002.png" in loaded["markdown"]
    assert "Second page requirement" in parser.select(loaded, {"kind": "page", "page": 2})[0]["text"]
    from docling_core.types.doc import DoclingDocument
    DoclingDocument.model_validate(loaded["document"])


def test_full_vector_pdf_does_not_automatically_render_pages(tmp_path, monkeypatch):
    import pypdfium2 as pdfium
    monkeypatch.setattr(pdfium.PdfPage, "render", lambda *a, **k: pytest.fail("Full ingestion rendered a page"))
    descriptor, loaded = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes())
    assert "three" in loaded["text"]
    assert loaded["manifest"]["page_preview"] is None
    assert not any(image["role"] == "page_preview" for image in descriptor["images"])


def test_failed_page_preview_keeps_text_and_native_locations(tmp_path, monkeypatch):
    import pypdfium2 as pdfium
    def fail(*args, **kwargs):
        raise RuntimeError("Synthetic renderer failure")
    monkeypatch.setattr(pdfium.PdfPage, "render", fail)
    descriptor, loaded = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes(), page_range=(2, 2))
    assert descriptor["parse_status"] == "partial"
    assert any(g["code"] == "PAGE_PREVIEW_UNAVAILABLE" for g in descriptor["gaps"])
    assert "Second page requirement" in parser.select(loaded, {"kind": "page", "page": 2})[0]["text"]
    assert loaded["manifest"]["page_preview"]["rendered_pages"] == []


@pytest.mark.parametrize("pages,size,expected", [(9, None, "PAGE_PREVIEW_PAGE_LIMIT"), (1, (50000, 50000), "PAGE_PREVIEW_PIXEL_LIMIT")])
def test_page_preview_page_and_pixel_budgets_reject_before_render(tmp_path, monkeypatch, pages, size, expected):
    import pypdfium2 as pdfium
    monkeypatch.setattr(pdfium.PdfPage, "render", lambda *a, **k: pytest.fail("An over-budget page was rendered"))
    descriptor, loaded = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes(pages, size), page_range=(1, pages))
    assert any(g["code"] == expected for g in descriptor["gaps"])
    assert loaded["manifest"]["page_preview"]["rendered_pages"] == []


def test_page_preview_png_budget_does_not_save_the_oversized_image(tmp_path, monkeypatch):
    monkeypatch.setattr(parser, "_PREVIEW_LIMITS", {**parser._PREVIEW_LIMITS, "max_png_bytes": 1})
    descriptor, loaded = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes(), page_range=(2, 2))
    assert any(g["code"] == "PAGE_PREVIEW_BYTE_LIMIT" for g in descriptor["gaps"])
    assert loaded["manifest"]["page_preview"]["png_bytes"] == 0
    assert not any(image["role"] == "page_preview" for image in descriptor["images"])


def test_unavailable_picture_bytes_remain_explicit_beside_a_page_preview(tmp_path, monkeypatch):
    original = parser._docling
    def unavailable_picture(*args):
        text, segments, images, gaps, pages = original(*args)
        gaps.append({"code": "IMAGE_UNAVAILABLE", "message": "Synthetic upstream bitmap bytes unavailable", "severity": "review"})
        images.append({"ref": None, "element_ref": "#/pictures/unavailable", "role": "document_picture"})
        return text, segments, images, gaps, pages
    monkeypatch.setattr(parser, "_docling", unavailable_picture)
    descriptor, _ = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes(), page_range=(2, 2))
    assert descriptor["parse_status"] == "partial"
    assert any(g["code"] == "IMAGE_UNAVAILABLE" for g in descriptor["gaps"])
    assert any(image["ref"] and image["role"] == "page_preview" for image in descriptor["images"])


def test_page_preview_asset_damage_is_rejected_without_rerender(tmp_path, monkeypatch):
    descriptor, _ = parse_bytes(tmp_path, ".pdf", data=vector_pdf_bytes(), page_range=(2, 2))
    preview = next(image for image in descriptor["images"] if image["role"] == "page_preview")
    path = tmp_path / preview["ref"]
    path.write_bytes(path.read_bytes() + b"changed")
    monkeypatch.setattr(parser, "_pdf_page_previews", lambda *a, **k: pytest.fail("load rendered a new preview"))
    with pytest.raises(ContractError, match="PARSE_MATERIAL_CHANGED"):
        parser.load(tmp_path, descriptor, expected_source_id=descriptor["source_id"], expected_sha256=descriptor["source_sha256"])


def test_scoped_elements_cannot_alias_a_different_full_document_element(tmp_path):
    data = vector_pdf_bytes()
    _, full = parse_bytes(tmp_path, ".pdf", data=data, output=".build/materials/references/full")
    _, scoped = parse_bytes(tmp_path, ".pdf", data=data, output=".build/materials/references/scoped", page_range=(2, 2))
    full_profile = full["manifest"]["profile"]
    assert "page_range" not in full_profile
    assert full_profile == parser._profile(parser._options(tmp_path, None), ".pdf", tmp_path)
    assert scoped["manifest"]["profile"]["page_range"] == [2, 2]
    element = next(row["locator"] for row in scoped["locators"]["segments"] if row["locator"]["kind"] == "element")
    assert parser.select(scoped, element)
    with pytest.raises(ContractError, match="INVALID_SOURCE_LOCATOR"):
        parser.select(full, element)


def test_native_sheet_range_and_slide_have_distinct_addresses(tmp_path):
    _, workbook = parse_bytes(tmp_path, ".xlsx", output=".build/materials/references/sheet-identity")
    cells = parser.select(workbook, {"kind": "cell_range", "sheet": "Config", "range": "B2"})
    assert [row["text"] for row in cells] == ["three"]
    with pytest.raises(ContractError, match="INVALID_SOURCE_LOCATOR"):
        parser.select(workbook, {"kind": "cell_range", "sheet": "Missing", "range": "B2"})
    _, slides = parse_bytes(tmp_path, ".pptx", output=".build/materials/references/slides-identity")
    assert "three" in parser.select(slides, {"kind": "slide", "slide": 1})[0]["text"]


def test_nondefault_element_requires_its_reproducible_parser_recipe(tmp_path):
    _, loaded = parse_bytes(tmp_path, ".docx", options={"threads": 2})
    locator = next(row["locator"] for row in loaded["locators"]["segments"] if row["locator"]["kind"] == "element")
    for recipe in (None, {}):
        with pytest.raises(ContractError, match="INVALID_SOURCE_LOCATOR"):
            parser.select(loaded, locator if recipe is None else {**locator, "parser": recipe})
    assert parser.select(loaded, {**locator, "parser": {"threads": 2}})


def test_bundle_damage_and_descriptor_substitution_never_reparse(tmp_path, monkeypatch):
    descriptor, _ = parse_bytes(tmp_path)
    monkeypatch.setattr(parser, "parse", lambda *a, **k: pytest.fail("load must never reparse"))
    with pytest.raises(ContractError, match="PARSE_SOURCE_MISMATCH"):
        parser.load(tmp_path, descriptor, expected_source_id="another", expected_sha256=descriptor["source_sha256"])
    (tmp_path / descriptor["text_ref"]).write_text("Changed material", encoding="utf-8")
    with pytest.raises(ContractError, match="PARSE_MATERIAL_CHANGED"):
        parser.load(tmp_path, descriptor, expected_source_id=descriptor["source_id"], expected_sha256=descriptor["source_sha256"])


@pytest.mark.parametrize("output", ["knowledge/parsed", ".build/parsed", ".build/materials/other", "artifacts/query/parsed", "artifacts/source-read/not-a-uuid/parsed", "artifacts/source-read/" + uuid.uuid4().hex + "/other", "../parsed"])
def test_parser_cannot_write_outside_operation_scopes(tmp_path, output):
    with pytest.raises(ContractError):
        parse_bytes(tmp_path, output=output)
    assert not (tmp_path / "knowledge").exists()


def test_each_read_gets_an_independent_directory_and_existing_outputs_are_not_reused(tmp_path):
    one, _ = parse_bytes(tmp_path, output=f"artifacts/source-read/{uuid.uuid4().hex}/parsed")
    two, _ = parse_bytes(tmp_path, output=f"artifacts/source-read/{uuid.uuid4().hex}/parsed")
    assert one["manifest_ref"] != two["manifest_ref"]
    assert one["source_sha256"] == two["source_sha256"]
    with pytest.raises(ContractError, match="PARSE_OUTPUT_EXISTS"):
        parse_bytes(tmp_path, output=str(Path(one["manifest_ref"]).parent))


@pytest.mark.parametrize("options", [{"unexpected": True}, {"threads": True}, {"threads": 33}, {"threads": 10 ** 1000}, {"pipeline": []}, {"timeout_seconds": float("nan")}, {"timeout_seconds": 10 ** 1000}, {"timeout_seconds": 3601}, {"images_scale": 10 ** 1000}, {"images_scale": 4.01}, {"max_input_bytes": 256 * 1024 * 1024 + 1}, {"max_output_bytes": 256 * 1024 * 1024 + 1}, {"models_path": "/tmp/models"}, {"models_path": "artifacts/other/models"}])
def test_options_are_explicit_and_model_paths_are_scoped(tmp_path, options):
    with pytest.raises(ContractError):
        parse_bytes(tmp_path, options=options)
    assert not (tmp_path / ".build").exists()


def test_missing_preinstalled_models_report_gap_without_network(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        pytest.fail("Parser attempted network access")
    monkeypatch.setattr(socket.socket, "connect", no_network)
    descriptor, _ = parse_bytes(tmp_path, ".pdf", options={"pipeline": "standard"})
    assert descriptor["parse_status"] == "partial"
    assert descriptor["gaps"][0]["code"] == "PARSER_MODELS_REQUIRED"


def test_empty_model_directory_is_not_a_download_permission(tmp_path, monkeypatch):
    (tmp_path / "artifacts/models/empty").mkdir(parents=True)
    monkeypatch.setattr(socket.socket, "connect", lambda *a, **k: pytest.fail("Missing model triggered network fallback"))
    descriptor, _ = parse_bytes(tmp_path, ".pdf", options={"pipeline": "standard", "models_path": "artifacts/models/empty"})
    assert descriptor["parse_status"] == "partial"
    assert descriptor["gaps"][0]["code"] == "PARSER_MODELS_REQUIRED"


def test_external_html_images_are_not_fetched_or_implicitly_read(tmp_path, monkeypatch):
    monkeypatch.setattr(socket.socket, "connect", lambda *a, **k: pytest.fail("External resource was fetched"))
    content = b'<html><body><p>Retry three times.</p><img src="https://example.invalid/secret.png"><img src="file:///etc/passwd"></body></html>'
    descriptor, loaded = parse_bytes(tmp_path, ".html", data=content)
    assert "three" in loaded["text"]
    assert len([g for g in descriptor["gaps"] if g["code"] == "EXTERNAL_IMAGE_NOT_READ"]) == 2


@pytest.mark.parametrize("suffix", [".md", ".txt"])
def test_existing_plain_text_source_bytes_are_not_rewritten(tmp_path, suffix):
    original = b"\xef\xbb\xbf# Requirement\r\n\r\nRetry three times.\r\n"
    descriptor, loaded = parse_bytes(tmp_path, suffix, data=original)
    assert descriptor["source_sha256"] == parser.sha256_bytes(original)
    decoded = original.decode("utf-8-sig")
    assert (tmp_path / descriptor["text_ref"]).read_bytes() == decoded.encode("utf-8")
    assert (tmp_path / descriptor["markdown_ref"]).read_bytes() == decoded.encode("utf-8")
    assert loaded["text"] == "# Requirement\n\nRetry three times.\n"
    assert len(loaded["text"].splitlines()) == 3
    segments = loaded["locators"]["segments"]
    assert [row["locator"]["start"] for row in segments] == [1, 2, 3]
    assert [row["text"] for row in segments] == decoded.splitlines()


@pytest.mark.parametrize("syntax", ["![figure](https://example.invalid/picture.png)", "![figure](../private.png)", "![figure][asset]\n\n[asset]: https://example.invalid/picture.png", '<img src="file:///etc/passwd">'])
def test_unread_native_markdown_images_cannot_make_renderer_fetch_resources(tmp_path, syntax):
    descriptor, loaded = parse_bytes(tmp_path, data=("Requirement.\n" + syntax).encode())
    assert "![" not in loaded["markdown"]
    assert "<img" not in loaded["markdown"]
    assert "Image not read" in loaded["markdown"]
    assert any(g["code"] == "EXTERNAL_IMAGE_NOT_READ" for g in descriptor["gaps"])


@pytest.mark.parametrize("image_uri", ["https://example.invalid/picture.png", "../private.png", "data:image/png;base64,AA=="])
def test_consistent_bundle_hashes_do_not_authorize_unregistered_markdown_images(tmp_path, image_uri):
    descriptor, _ = parse_bytes(tmp_path)
    markdown_path = tmp_path / descriptor["markdown_ref"]
    markdown_path.write_text("![Image](" + image_uri + ")", encoding="utf-8")
    manifest_path = tmp_path / descriptor["manifest_ref"]
    manifest = parser.read_json(manifest_path)
    for row in manifest["files"]:
        if row["ref"] == descriptor["markdown_ref"]:
            raw = markdown_path.read_bytes()
            row.update(bytes=len(raw), sha256=parser.sha256_bytes(raw))
    manifest_path.write_bytes(parser.canonical_bytes(manifest))
    descriptor["manifest_sha256"] = parser.sha256_bytes(manifest_path.read_bytes())
    with pytest.raises(ContractError):
        parser.load(tmp_path, descriptor, expected_source_id=descriptor["source_id"], expected_sha256=descriptor["source_sha256"])


@pytest.mark.parametrize("read_id", [str(uuid.uuid4()), uuid.uuid4().hex.upper()])
def test_read_scope_uuid_has_one_canonical_spelling(tmp_path, read_id):
    with pytest.raises(ContractError, match="PARSE_OUTPUT_UNSAFE"):
        parse_bytes(tmp_path, output=f"artifacts/source-read/{read_id}/parsed")


def test_limits_reject_before_silently_truncating_material(tmp_path):
    with pytest.raises(ContractError, match="PARSE_INPUT_LIMIT"):
        parse_bytes(tmp_path, options={"max_input_bytes": 1})
    with pytest.raises(ContractError, match="PARSE_OUTPUT_LIMIT"):
        parse_bytes(tmp_path, options={"max_output_bytes": 1})


def test_linked_output_directory_is_rejected(tmp_path):
    elsewhere = tmp_path / "unrelated"
    elsewhere.mkdir()
    materials = tmp_path / ".build/materials"
    materials.mkdir(parents=True)
    try:
        (materials / "parsed").symlink_to(elsewhere, target_is_directory=True)
    except OSError:
        pytest.skip("The host does not allow synthetic symlinks")
    with pytest.raises(ContractError):
        parse_bytes(tmp_path)
    assert not list(elsewhere.iterdir())


def test_internal_link_cannot_redirect_new_materials_into_readonly_originals(tmp_path):
    originals = tmp_path / "docs-raw"
    originals.mkdir()
    sentinel = originals / "original.txt"
    sentinel.write_bytes(b"Read-only original")
    materials = tmp_path / ".build/materials"
    materials.mkdir(parents=True)
    try:
        (materials / "references").symlink_to(originals, target_is_directory=True)
    except OSError:
        pytest.skip("The host does not allow synthetic symlinks")
    with pytest.raises(ContractError):
        parse_bytes(tmp_path, output=".build/materials/references/redirected")
    assert list(originals.iterdir()) == [sentinel]
    assert sentinel.read_bytes() == b"Read-only original"


def test_load_rejects_internal_symlink_even_if_target_has_matching_bytes(tmp_path):
    descriptor, _ = parse_bytes(tmp_path)
    text = tmp_path / descriptor["text_ref"]
    other = tmp_path / "same-bytes.txt"
    other.write_bytes(text.read_bytes())
    text.unlink()
    try:
        text.symlink_to(other)
    except OSError:
        pytest.skip("The host does not allow synthetic symlinks")
    with pytest.raises(ContractError):
        parser.load(tmp_path, descriptor, expected_source_id=descriptor["source_id"], expected_sha256=descriptor["source_sha256"])
