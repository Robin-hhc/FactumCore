"""Author/reviewer integration uses real synthetic documents and exact P1 bundles."""

from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import sys
from threading import Event

import pytest
import yaml


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "support"))
from contracts import module, write_plan


def library(tmp_path):
    root = tmp_path / "library"
    for directory in ("docs-raw", "knowledge", "governance", "artifacts"):
        (root / directory).mkdir(parents=True)
    return root


def candidate(root, identity, reference):
    path = root / ".build/cards" / (identity + ".md")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("---\n" + yaml.safe_dump({"id": identity, "doc_ref": [reference], "code_ref": []})
                    + "---\n# Documented requirement\n\nA design requirement; applicability needs review.\n", encoding="utf-8")
    write_plan(root, [(identity, "knowledge/" + identity + ".md")])
    return path


def scan(root, name, **extra):
    task, status = module("ingest").scan(root, {"source": "docs-raw/" + name, **extra})
    assert status == "prepared", task
    return task


def test_p4_reuses_the_fixed_p1_bundle_and_p5_cleans_only_ingest(tmp_path, monkeypatch):
    root = library(tmp_path)
    original = root / "docs-raw/design.md"
    original.write_text("# Design\n\nThe bounded design requirement.\n", encoding="utf-8")
    before = original.read_bytes()
    task = scan(root, original.name)
    parsed = task["source"]["parsed"]
    assert parsed["manifest_ref"].startswith(".build/materials/parsed/")
    assert (root / ".build/materials/current.txt").read_bytes() == (root / parsed["text_ref"]).read_bytes()
    independent = root / "artifacts/source-read/independent-marker"
    independent.parent.mkdir(parents=True)
    independent.write_text("another reader", encoding="utf-8")
    candidate(root, "requirement", {"doc_id": task["source"]["source_id"], "line_start": 1, "line_end": 3})
    monkeypatch.setattr(module("document_parser"), "parse", lambda *a, **k: pytest.fail("P4 must not reconvert P1"))
    checked, status = module("ingest").check(root, {})
    assert status == "passed", checked
    committed, status = module("ingest").commit(root, {})
    assert status == "completed", committed
    assert not (root / ".build").exists()
    assert independent.read_text() == "another reader"
    assert original.read_bytes() == before


@pytest.mark.parametrize("source_kind", ["raw", "code"])
def test_txt_original_with_bom_and_crlf_survives_p4_and_p5(tmp_path, source_kind):
    root = library(tmp_path)
    original_bytes = b"\xef\xbb\xbfA documented requirement.\r\nKeep its original bytes.\r\n"
    if source_kind == "raw":
        original = root / "docs-raw/design.txt"
        original.write_bytes(original_bytes)
        task = scan(root, original.name)
    else:
        original = tmp_path / "main/wifi/design.txt"
        original.parent.mkdir(parents=True)
        original.write_bytes(original_bytes)
        (tmp_path / "config.yml").write_text(yaml.safe_dump({
            "schema_version": "2.0", "workspace_id": "txt-copy-test", "knowledge_roots": {"wifi": "library"},
            "code_root": "main", "modules": [{"path": "wifi"}],
        }), encoding="utf-8")
        module("library").sources(root, "register", {
            "kind": "code", "path": "main/wifi/design.txt", "source_id": "code-text", "encoding": "utf-8"})
        task, status = module("ingest").scan(root, {"source": "code-text"})
        assert status == "prepared", task
    source = task["source"]
    assert (root / source["original_ref"]).read_bytes() == original_bytes
    assert source["text_ref"] != source["original_ref"]
    assert not (root / source["text_ref"]).read_bytes().startswith(b"\xef\xbb\xbf")
    assert source["hash"] == module("ingest").sha(original_bytes)
    write_plan(root, [])
    checked, status = module("ingest").check(root, {})
    assert status == "passed", checked
    assert (root / source["original_ref"]).read_bytes() == original_bytes
    committed, status = module("ingest").commit(root, {"reason": "Reviewed source adds no additional knowledge."})
    assert status == "completed" and committed["processed_hash"] == source["hash"], committed
    assert not (root / ".build").exists()
    assert original.read_bytes() == original_bytes


@pytest.mark.parametrize("tamper", ["original", "text"])
def test_p4_rejects_damaged_p1_material_even_for_a_no_card_result(tmp_path, tamper):
    root = library(tmp_path)
    (root / "docs-raw/design.md").write_text("# Design\n\nA requirement.\n", encoding="utf-8")
    task = scan(root, "design.md")
    write_plan(root, [])
    reference = task["source"]["original_ref"] if tamper == "original" else task["source"]["parsed"]["text_ref"]
    (root / reference).write_text("changed bytes", encoding="utf-8")
    checked, status = module("ingest").check(root, {})
    assert status == "failed", checked
    assert checked["findings"][0]["code"] in {"SOURCE_CHANGED", "PARSE_MATERIAL_CHANGED"}


def test_explicit_reprocess_preserves_success_facts_until_p5(tmp_path):
    root = library(tmp_path)
    (root / "docs-raw/design.md").write_text("# Design\n\nA requirement.\n", encoding="utf-8")
    task = scan(root, "design.md")
    identity = task["source"]["source_id"]
    module("ingest").commit(root, {"reason": "Reviewed source has no additional knowledge."})
    registry = (root / "governance/sources.lock.yaml").read_bytes()
    assert module("ingest").scan(root, {})[1] == "unchanged"
    repeated, status = module("ingest").scan(root, {"source": identity, "reprocess": True})
    assert status == "prepared" and repeated["source"]["change"] == "reprocess"
    assert repeated["source"]["hash"] == task["source"]["hash"]
    assert (root / "governance/sources.lock.yaml").read_bytes() == registry
    write_plan(root, [])
    assert module("ingest").check(root, {})[1] == "passed"


@pytest.mark.parametrize("payload", [{"reprocess": True}, {"source": "docs-raw", "reprocess": True},
                                      {"source": "unregistered", "reprocess": True}, {"reprocess": "true"}])
def test_invalid_reprocess_does_not_discard_the_current_task(tmp_path, payload):
    root = library(tmp_path)
    marker = root / ".build/task.json"
    marker.parent.mkdir()
    marker.write_bytes(b'{"other_task":true}')
    with pytest.raises(module("validation").ContractError):
        module("ingest").scan(root, payload)
    assert marker.read_bytes() == b'{"other_task":true}'


@pytest.mark.parametrize("kind", ["xlsx", "pptx"])
def test_new_native_locators_are_checked_against_real_document_locations(tmp_path, kind):
    root = library(tmp_path)
    path = root / "docs-raw" / ("design." + kind)
    if kind == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Requirements"
        sheet.append(["Name", "Value"])
        sheet.append(["budget", 8])
        book.save(path)
        good = {"kind": "cell_range", "sheet": "Requirements", "range": "A1:B2"}
        bad = {**good, "sheet": "Missing"}
    else:
        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Design requirement"
        slide.placeholders[1].text = "The budget is eight milliseconds by design."
        deck.save(path)
        good, bad = {"kind": "slide", "slide": 1}, {"kind": "slide", "slide": 2}
    task = scan(root, path.name)
    assert task["source"]["parsed"]["parse_status"] == "parsed", task
    identity = task["source"]["source_id"]
    candidate(root, "requirement", {"doc_id": identity, "locator": good})
    assert module("ingest").check(root, {})[1] == "passed"
    candidate(root, "requirement", {"doc_id": identity, "locator": bad})
    checked, status = module("ingest").check(root, {})
    assert status == "failed" and checked["findings"][0]["code"] == "INVALID_SOURCE_LOCATOR", checked
    candidate(root, "requirement", {"doc_id": identity, "line_start": 1, "line_end": 1})
    checked, status = module("ingest").check(root, {})
    assert status == "failed" and checked["findings"][0]["code"] == "INVALID_SOURCE_LOCATOR", checked


def test_element_locator_binds_both_source_bytes_and_parser_profile(tmp_path):
    from docx import Document

    root = library(tmp_path)
    path = root / "docs-raw/design.docx"
    document = Document()
    document.add_paragraph("A requirement with a native document element.")
    document.save(path)
    task = scan(root, path.name)
    parsed = task["source"]["parsed"]
    locators = json.loads((root / parsed["locators_ref"]).read_bytes())
    locator = next(row["locator"] for row in locators["segments"] if row["locator"]["kind"] == "element")
    identity = task["source"]["source_id"]
    candidate(root, "requirement", {"doc_id": identity, "locator": locator})
    assert module("ingest").check(root, {})[1] == "passed"
    for field in ("source_sha256", "profile_sha256"):
        candidate(root, "requirement", {"doc_id": identity, "locator": {**locator, field: "0" * 64}})
        checked, status = module("ingest").check(root, {})
        assert status == "failed" and checked["findings"][0]["code"] == "INVALID_SOURCE_LOCATOR", checked


def test_nondefault_element_profile_is_reproducible_after_p5_cleans_the_bundle(tmp_path):
    from docx import Document
    from jsonschema import validate

    root = library(tmp_path)
    path = root / "docs-raw/design.docx"
    document = Document()
    document.add_paragraph("A requirement with an explicit parser profile.")
    document.save(path)
    task = scan(root, path.name, parser={"threads": 2})
    parsed = task["source"]["parsed"]
    loaded = module("document_parser").load(root, parsed, expected_source_id=task["source"]["source_id"],
                                             expected_sha256=task["source"]["hash"])
    locator = next(row["locator"] for row in loaded["locators"]["segments"] if row["locator"]["kind"] == "element")
    locator = {**locator, "parser": loaded["manifest"]["profile"]["options"]}
    schema = json.loads((Path(__file__).resolve().parents[2] / "governance/schemas/frontmatter.schema.json").read_bytes())
    validate(locator, schema["$defs"]["document_locator"])
    for missing_or_wrong in ({}, {"threads": 3}):
        candidate(root, "requirement", {"doc_id": task["source"]["source_id"],
                                         "locator": {**locator, "parser": missing_or_wrong}})
        checked, status = module("ingest").check(root, {})
        assert status == "failed" and checked["findings"][0]["code"] == "INVALID_SOURCE_LOCATOR", checked
    candidate(root, "requirement", {"doc_id": task["source"]["source_id"], "locator": locator})
    checked, status = module("ingest").check(root, {})
    assert status == "passed", checked
    committed, status = module("ingest").commit(root, {})
    assert status == "completed" and not (root / ".build").exists(), committed
    checked, status = module("ingest").check(root, {"formal": True})
    assert status == "passed", checked
    assert not (root / ".build").exists()
    formal = root / "knowledge/requirement.md"
    text = formal.read_text(encoding="utf-8")
    header, body = text.split("---\n", 2)[1:]
    metadata = yaml.safe_load(header)
    del metadata["doc_ref"][0]["locator"]["parser"]
    formal.write_text("---\n" + yaml.safe_dump(metadata) + "---\n" + body, encoding="utf-8")
    checked, status = module("ingest").check(root, {"formal": True})
    assert status == "failed" and checked["findings"][0]["code"] == "INVALID_SOURCE_LOCATOR", checked


@pytest.mark.parametrize("secondary_id", ["doc-secondary", "doc:secondary"])
def test_secondary_source_is_parsed_once_and_formal_lint_keeps_author_work_intact(tmp_path, monkeypatch, secondary_id):
    from pptx import Presentation

    root = library(tmp_path)
    (root / "docs-raw/design.md").write_text("# Design\n\nA requirement.\n", encoding="utf-8")
    task = scan(root, "design.md")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Secondary evidence"
    slide.placeholders[1].text = "A requirement supported by this source."
    deck.save(root / "docs-raw/secondary.pptx")
    registered, _ = module("library").sources(root, "register", {
        "path": "docs-raw/secondary.pptx", "source_id": secondary_id})
    identity = registered["source"]["source_id"]
    path = candidate(root, "requirement", {"doc_id": identity, "locator": {"kind": "slide", "slide": 1}})
    parser = module("document_parser")
    parse = parser.parse
    calls = []

    def count_parse(*args, **kwargs):
        calls.append(args[1]["source_id"])
        return parse(*args, **kwargs)

    monkeypatch.setattr(parser, "parse", count_parse)
    first, status = module("ingest").check(root, {})
    assert status == "passed", first
    second, status = module("ingest").check(root, {})
    assert status == "passed", second
    assert calls == [identity]
    assert identity in first["parsed_sources"] and task["source"]["source_id"] not in first["parsed_sources"]
    formal = root / "knowledge/requirement.md"
    formal.write_bytes(path.read_bytes())
    before = {item.relative_to(root / ".build").as_posix(): item.read_bytes()
              for item in (root / ".build").rglob("*") if item.is_file()}
    checked, status = module("ingest").check(root, {"formal": True})
    assert status == "passed", checked
    assert {item.relative_to(root / ".build").as_posix(): item.read_bytes()
            for item in (root / ".build").rglob("*") if item.is_file()} == before
    assert not [item for item in (root / "artifacts/source-read").glob("*") if item.is_dir() and not item.name.startswith(".")]


def test_secondary_parse_index_rejects_an_internal_link_before_conversion(tmp_path, monkeypatch):
    from pptx import Presentation

    root = library(tmp_path)
    (root / "docs-raw/design.md").write_text("# Design\n\nA requirement.\n", encoding="utf-8")
    scan(root, "design.md")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Secondary source"
    deck.save(root / "docs-raw/secondary.pptx")
    registered, _ = module("library").sources(root, "register", {"path": "docs-raw/secondary.pptx"})
    candidate(root, "requirement", {"doc_id": registered["source"]["source_id"], "locator": {"kind": "slide", "slide": 1}})
    original = root / "docs-raw/keep.txt"
    original.write_bytes(b"{}")
    index = root / ".build/materials/reference-parses.json"
    index.symlink_to(original)

    def forbid_conversion(*args, **kwargs):
        pytest.fail("An unsafe index must be rejected before document conversion")

    monkeypatch.setattr(module("document_parser"), "parse", forbid_conversion)
    checked, status = module("ingest").check(root, {})
    assert status == "failed" and checked["findings"][0]["code"] == "WORKSPACE_LINKED_PATH", checked
    assert original.read_bytes() == b"{}" and index.is_symlink()
    assert not (root / ".build/materials/references").exists()


def test_secondary_parse_index_is_checked_again_before_writing(tmp_path, monkeypatch):
    root = library(tmp_path)
    original = root / "docs-raw/keep.txt"
    original.write_bytes(b"{}")
    source = {"source_id": "doc-secondary", "path": "docs-raw/secondary.md"}
    data = b"# Secondary source\nA real native text conversion.\n"
    (root / source["path"]).write_bytes(data)
    index = root / ".build/materials/reference-parses.json"
    parser = module("document_parser")
    parse = parser.parse

    def introduce_link_after_conversion(*args, **kwargs):
        parsed = parse(*args, **kwargs)
        index.symlink_to(original)
        return parsed

    monkeypatch.setattr(parser, "parse", introduce_link_after_conversion)
    context = {"formal": False, "loaded": {}, "descriptors": {}}
    with pytest.raises(module("validation").ContractError) as error:
        module("ingest")._parsed_document(root, source, data, context)
    assert error.value.code == "WORKSPACE_LINKED_PATH"
    assert original.read_bytes() == b"{}" and index.is_symlink()


@pytest.mark.parametrize("locator", [
    {"kind": "slide", "slide": 2}, {"kind": "cell_range", "sheet": "Sheet 1", "range": "A1:D20"},
    {"kind": "element", "ref": "#/tables/0", "source_sha256": "a" * 64, "profile_sha256": "b" * 64},
])
def test_new_native_locations_are_accepted_by_both_schema_and_card_parser(locator):
    from jsonschema import validate

    schema = json.loads((Path(__file__).resolve().parents[2] / "governance/schemas/frontmatter.schema.json").read_bytes())
    validate(locator, schema["$defs"]["document_locator"])
    module("cards_v3")._document_locator(locator, "knowledge/example.md")


@pytest.mark.parametrize("address,alter,expected", [
    ({"kind": "paragraph", "paragraph": 1}, None, "passed"),
    ({"kind": "paragraph", "paragraph": 2}, None, "INVALID_SOURCE_LOCATOR"),
    ({"kind": "paragraph", "paragraph": 1}, "extractor", "SOURCE_ADAPTER_UNAVAILABLE"),
    ({"kind": "lines", "start": 1, "end": 1}, None, "passed"),
    ({"kind": "lines", "start": 2, "end": 2}, None, "SOURCE_RANGE_INVALID"),
    ({"kind": "lines", "start": 1, "end": 1}, "text_sha256", "SOURCE_EXTRACTION_UNAVAILABLE"),
    ({"kind": "paragraph", "paragraph": 1}, "source_id", "SOURCE_SNAPSHOT_MISMATCH"),
])
def test_real_retained_snapshot_uses_exact_old_document_and_adapter(tmp_path, address, alter, expected):
    from docx import Document

    root = library(tmp_path)
    path = root / "docs-raw/design.docx"
    document = Document()
    document.add_paragraph("The retained original has exactly one paragraph.")
    document.save(path)
    old_bytes = path.read_bytes()
    evidence = root / "logs/evidence/old.bin"
    evidence.parent.mkdir(parents=True)
    evidence.write_bytes(old_bytes)
    module("library").sources(root, "register", {"path": "docs-raw/design.docx", "source_id": "doc-design"})
    text, addresses, extractor, _ = module("source_resources").extract_document(old_bytes, ".docx")
    extraction = {"extractor": extractor, "text_sha256": module("ingest").sha(text.encode()),
                  "locator_sha256": module("validation").sha256_bytes(module("validation").canonical_bytes({"locators": addresses}))}
    if alter == "source_id":
        (root / "docs-raw/other.md").write_text("# Other source\n\nA different document.\n", encoding="utf-8")
        module("library").sources(root, "register", {"path": "docs-raw/other.md", "source_id": "doc-other"})
    elif alter is not None:
        extraction[alter] = "unavailable-adapter" if alter == "extractor" else "0" * 64
    registry_path = root / "governance/sources.lock.yaml"
    registry = yaml.safe_load(registry_path.read_bytes())
    registry["snapshots"].append({"snapshot_id": "old-design", "source_id": "doc-other" if alter == "source_id" else "doc-design", "format": ".docx",
                                  "sha256": module("ingest").sha(old_bytes), "bytes": len(old_bytes),
                                  "recovery": {"method": "evidence_object", "path": "logs/evidence/old.bin"},
                                  "extraction": extraction})
    registry_path.write_text(yaml.safe_dump(registry, sort_keys=False), encoding="utf-8")
    document.add_paragraph("Only the current source has paragraph two.")
    document.save(path)
    reference = {"doc_id": "doc-design", "snapshot_id": "old-design"}
    if address["kind"] == "lines":
        reference.update({"line_start": address["start"], "line_end": address["end"]})
    else:
        reference["locator"] = address
    metadata = {"id": "history", "doc_ref": [reference], "code_ref": []}
    formal = root / "knowledge/history.md"
    formal.write_text("---\n" + yaml.safe_dump(metadata) + "---\n# Historical requirement\n\nHistorical evidence.\n", encoding="utf-8")
    original, card, lock = path.read_bytes(), formal.read_bytes(), registry_path.read_bytes()
    checked, status = module("ingest").check(root, {"formal": True})
    if expected == "passed":
        assert status == "passed", checked
    else:
        assert status == "failed" and checked["findings"][0]["code"] == expected, checked
    assert path.read_bytes() == original and formal.read_bytes() == card and registry_path.read_bytes() == lock
    assert not (root / ".build").exists()


def test_conversion_does_not_hold_the_catalog_transaction_against_queries(tmp_path, monkeypatch):
    root = library(tmp_path)
    (root / "docs-raw/design.md").write_text("# Design\n\nA requirement.\n", encoding="utf-8")
    parser = module("document_parser")
    parse = parser.parse
    entered, proceed = Event(), Event()

    def slow_parse(*args, **kwargs):
        entered.set()
        assert proceed.wait(15)
        return parse(*args, **kwargs)

    monkeypatch.setattr(parser, "parse", slow_parse)
    with ThreadPoolExecutor(max_workers=1) as workers:
        future = workers.submit(module("ingest").scan, root, {})
        try:
            assert entered.wait(10)
            result, status = module("library").query(root, {"query": "requirement"})
            assert status == "searched" and result["matches"] == []
        finally:
            proceed.set()
        assert future.result()[1] == "prepared"


@pytest.mark.parametrize("locator", [
    {"kind": "slide", "slide": True}, {"kind": "slide", "slide": 0},
    {"kind": "cell_range", "sheet": "S", "range": "B2:A1"},
    {"kind": "cell_range", "sheet": "S", "range": "XFE1"},
    {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64},
    {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64,
     "profile_sha256": "b" * 64, "parser": {"threads": True}},
    {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64,
     "profile_sha256": "b" * 64, "parser": {"models_path": "/machine/private/models"}},
    {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64,
     "profile_sha256": "b" * 64, "parser": {"pipeline": []}},
    {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64,
     "profile_sha256": "b" * 64, "parser": {"max_input_bytes": 256 * 1024 * 1024 + 1}},
    {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64,
     "profile_sha256": "b" * 64, "parser": {"max_output_bytes": 256 * 1024 * 1024 + 1}},
])
def test_new_locator_shapes_reject_ambiguous_or_unbound_addresses(locator):
    with pytest.raises(module("validation").ContractError):
        module("cards_v3")._document_locator(locator, "knowledge/example.md")


@pytest.mark.parametrize("field,value", [("threads", 33), ("timeout_seconds", 3601), ("images_scale", 4.1)])
def test_element_resource_limits_are_enforced_by_schema_and_candidate_validation(field, value):
    from jsonschema import ValidationError, validate

    locator = {"kind": "element", "ref": "#/texts/0", "source_sha256": "a" * 64,
               "profile_sha256": "b" * 64, "parser": {field: value}}
    with pytest.raises(module("validation").ContractError):
        module("cards_v3")._document_locator(locator, "knowledge/example.md")
    schema = json.loads((Path(__file__).resolve().parents[2] / "governance/schemas/frontmatter.schema.json").read_bytes())
    with pytest.raises(ValidationError):
        validate(locator, schema["$defs"]["document_locator"])
