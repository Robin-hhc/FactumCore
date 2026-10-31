"""Public checks preserve legacy document addresses without rewriting old cards."""

from base64 import b64decode
from pathlib import Path
import sys

from docx import Document
import pytest
import yaml


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "support"))
from contracts import module, write_plan


def test_encrypted_pdf_is_rejected_without_optional_crypto():
    # Fixed AES-256 blank PDF generated once with pypdf; no crypto is needed to
    # construct this fixture in the runtime under test.
    data = b64decode(
        b"JVBERi0xLjMKJeLjz9MKMSAwIG9iago8PAovUHJvZHVjZXIgPDQ3NDhjYjhiZGY2YWNlMDg0YTdkZmEy"
        b"OGI1NTRmMTljYzAxY2JjNDRmMmVlZWI5Mzc4Y2U1ZTMzYzdhZDFkOWM+Cj4+CmVuZG9iagoyIDAgb2Jq"
        b"Cjw8Ci9UeXBlIC9QYWdlcwovQ291bnQgMQovS2lkcyBbIDQgMCBSIF0KPj4KZW5kb2JqCjMgMCBvYmoK"
        b"PDwKL1R5cGUgL0NhdGFsb2cKL1BhZ2VzIDIgMCBSCj4+CmVuZG9iago0IDAgb2JqCjw8Ci9UeXBlIC9Q"
        b"YWdlCi9SZXNvdXJjZXMgPDwKPj4KL01lZGlhQm94IFsgMC4wIDAuMCAxMDAgMTAwIF0KL1BhcmVudCAy"
        b"IDAgUgo+PgplbmRvYmoKNSAwIG9iago8PAovViA1Ci9SIDYKL0xlbmd0aCAyNTYKL1AgNDI5NDk2NzI5"
        b"MgovRmlsdGVyIC9TdGFuZGFyZAovTyA8ZGI3OGU1ZjM1OWNmZWZhYjgyMGFkMzZkNTZiMTUyZGI0Y2My"
        b"MWEwYTUyMTQ1ZTAxY2JhNDAwZjAwNjE0MGZiYmNkOTE5YzEzODhlZmYzYmRmNzRjYjc0YWM2OGMwMWU1"
        b"PgovVSA8ZDA4OGZmOTdmNjhiZTcwZmJiYmVjOGQ5YTg1MmExNzIyZTYzOWI5OGNlOTMwYmI4ZTRlNWIw"
        b"Nzk1OTQzMDI5N2ZmN2YyMGY5Yzk2YWU4MTY4NGQ1NDRmNWMyNmQ1ZmRkPgovQ0YgPDwKL1N0ZENGIDw8"
        b"Ci9BdXRoRXZlbnQgL0RvY09wZW4KL0NGTSAvQUVTVjMKL0xlbmd0aCAzMgo+Pgo+PgovU3RtRiAvU3Rk"
        b"Q0YKL1N0ckYgL1N0ZENGCi9PRSA8OWI4NjJjZjcwNmZiOGE2NzI4Y2EwYmQ5ODYzNTYyZWQyNjUzY2Ez"
        b"ZDA3YzYxYjkxZjFlODhjNzBkOTc5MDQzNj4KL1VFIDxiNDZiM2VjNGY0ZWQ3NWRkNmIxZTZiZGU2MWE0"
        b"NWZkZDg4MGM5YTJhZGMzNjQyZTBhMzU1ZDFiYzA5OTk0MGRjPgovUGVybXMgPDY1OGVhN2Q2M2UxYjgz"
        b"MzAwODJlYTRjODJhMzhjYzc0Pgo+PgplbmRvYmoKeHJlZgowIDYKMDAwMDAwMDAwMCA2NTUzNSBmIAow"
        b"MDAwMDAwMDE1IDAwMDAwIG4gCjAwMDAwMDAxMTMgMDAwMDAgbiAKMDAwMDAwMDE3MiAwMDAwMCBuIAow"
        b"MDAwMDAwMjIxIDAwMDAwIG4gCjAwMDAwMDAzMTUgMDAwMDAgbiAKdHJhaWxlcgo8PAovU2l6ZSA2Ci9S"
        b"b290IDMgMCBSCi9JbmZvIDEgMCBSCi9JRCBbIDwzMDM5MzUzMzMxNjY2MzM4NjEzNzMwMzE2MzM3MzA2"
        b"MjM3NjMzNzMyMzMzMDMwMzYzODM1MzYzMjM2Mzg2NDYzPiA8MzAzOTM1MzMzMTY2NjMzODYxMzczMDMx"
        b"NjMzNzMwNjIzNzYzMzczMjMzMzAzMDM2MzgzNTM2MzIzNjM4NjQ2Mz4gXQovRW5jcnlwdCA1IDAgUgo+"
        b"PgpzdGFydHhyZWYKODcwCiUlRU9GCg=="
    )
    with pytest.raises(module("validation").ContractError, match="ENCRYPTED_DOCUMENT"):
        module("source_resources").extract_document(data, ".pdf")


@pytest.mark.parametrize("paragraph, expected", [(2, "passed"), (3, "failed")])
@pytest.mark.parametrize("card_format", ["formal-old", "concise-candidate"])
def test_legacy_docx_paragraph_is_checked_against_native_document(tmp_path, paragraph, expected, card_format):
    root = tmp_path / "library"
    for directory in ("docs-raw", "knowledge", "governance", "artifacts"):
        (root / directory).mkdir(parents=True)
    original = root / "docs-raw/design.docx"
    document = Document()
    document.add_paragraph("First documented requirement.")
    document.add_paragraph("Second documented requirement.")
    document.save(original)
    original_bytes = original.read_bytes()
    registered, status = module("library").sources(root, "register", {
        "kind": "raw", "path": "docs-raw/design.docx", "source_id": "doc-design",
    })
    assert status == "registered", registered
    metadata = {
        "schema_version": "2.0", "id": "requirement", "profile": "entity", "title": "Requirement",
        "summary": "A documented requirement", "lifecycle": "stable", "owner": "synthetic",
        "review_mode": "manual_document", "scope": {"applicability": "unknown", "contexts": []},
        "aliases": [], "relations": [], "sources": [{
            "evidence_ref": "doc", "source_id": "doc-design", "snapshot_id": "old-doc",
            "locator": {"kind": "paragraph", "paragraph": paragraph},
        }], "claims": [{
            "id": "answer", "text_ref": "#answer", "usage": "guidance", "evidence_kind": "design",
            "source_refs": ["doc"], "context_ids": [], "verification_requirements": [],
        }],
    }
    old_bytes = ("---\n" + yaml.safe_dump(metadata, sort_keys=False)
                 + '---\n# Requirement\n\n<a id="answer"></a>\nSecond documented requirement.\n').encode()
    formal = root / "knowledge/requirement.md"
    formal.write_bytes(old_bytes)
    if card_format == "formal-old":
        request = {"formal": True}
        candidate = None
    else:
        task, status = module("ingest").scan(root, {"source": "doc-design"})
        assert status == "prepared", task
        header = {"id": "requirement", "doc_ref": [{
            "doc_id": "doc-design", "locator": {"kind": "paragraph", "paragraph": paragraph},
            "snapshot_id": "old-doc", "evidence_ref": "doc",
        }], "code_ref": []}
        candidate = root / ".build/cards/requirement.md"
        candidate.write_text("---\n" + yaml.safe_dump(header, sort_keys=False)
                             + '---\n# Requirement\n\n<a id="answer"></a>\n'
                             + "Second documented requirement; applicability remains unknown.\n", encoding="utf-8")
        candidate_bytes = candidate.read_bytes()
        write_plan(root, [("requirement", "knowledge/requirement.md")],
                   body="Keep the old native document locator and anchor when adopting the concise header.")
        request = {}
    checked, status = module("ingest").check(root, request)

    assert status == expected, checked
    if expected == "failed":
        assert any(item["code"] == "INVALID_SOURCE_LOCATOR" for item in checked["findings"])
    else:
        assert checked["findings"] == []
    assert formal.read_bytes() == old_bytes
    if candidate is not None:
        assert candidate.read_bytes() == candidate_bytes
    assert original.read_bytes() == original_bytes


def write_native_document(path):
    """Create real readable DOCX/PDF bytes; no extractor is mocked."""
    if path.suffix == ".docx":
        document = Document()
        document.add_paragraph("DesignSection")
        document.add_paragraph("A documented requirement.")
        document.save(path)
    else:
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=180)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 24 140 Td (DesignSection) Tj 0 -18 Td (A documented requirement.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        with path.open("wb") as stream:
            writer.write(stream)


@pytest.mark.parametrize("suffix", [".docx", ".pdf"])
@pytest.mark.parametrize("locator_kind", ["lines", "section"])
@pytest.mark.parametrize("valid", [True, False], ids=["valid", "invalid"])
def test_legacy_pdf_docx_text_locators_use_text_checks(tmp_path, monkeypatch, suffix, locator_kind, valid):
    root = tmp_path / "library"
    for directory in ("docs-raw", "knowledge", "governance", "artifacts"):
        (root / directory).mkdir(parents=True)
    original = root / "docs-raw" / ("design" + suffix)
    write_native_document(original)
    original_bytes = original.read_bytes()
    registered, status = module("library").sources(root, "register", {
        "kind": "raw", "path": "docs-raw/" + original.name, "source_id": "doc-design",
    })
    assert status == "registered", registered
    if locator_kind == "lines":
        locator = {"kind": "lines", "start": 1, "end": 1 if valid else 999}
        expected_error = "SOURCE_RANGE_INVALID"
    else:
        locator = {"kind": "section", "section": "DesignSection" if valid else "MissingSection"}
        expected_error = "SOURCE_SECTION_MISSING"
    metadata = {
        "schema_version": "2.0", "id": "legacy", "profile": "rule", "title": "Legacy requirement",
        "summary": "A documented requirement", "lifecycle": "stable", "owner": "synthetic",
        "review_mode": "manual_document", "scope": {"applicability": "unknown", "contexts": []},
        "aliases": [], "relations": [], "sources": [{
            "evidence_ref": "doc", "source_id": "doc-design", "snapshot_id": "old-doc", "locator": locator,
        }], "claims": [{
            "id": "answer", "text_ref": "#answer", "usage": "guidance", "evidence_kind": "design",
            "source_refs": ["doc"], "context_ids": [], "verification_requirements": [],
        }],
    }
    old_bytes = ("---\n" + yaml.safe_dump(metadata, sort_keys=False)
                 + '---\n# Legacy requirement\n\n<a id="answer"></a>\nA documented requirement.\n').encode()
    formal = root / "knowledge/legacy.md"
    formal.write_bytes(old_bytes)

    def forbid_native_dispatch(*args, **kwargs):
        raise AssertionError("Legacy lines/section must use text checks for DOCX/PDF")

    monkeypatch.setattr(module("source_resources"), "document_segments", forbid_native_dispatch)
    checked, status = module("ingest").check(root, {"formal": True})

    assert status == ("passed" if valid else "failed"), checked
    if valid:
        assert checked["findings"] == []
    else:
        assert any(item["code"] == expected_error for item in checked["findings"])
    assert formal.read_bytes() == old_bytes
    assert original.read_bytes() == original_bytes
