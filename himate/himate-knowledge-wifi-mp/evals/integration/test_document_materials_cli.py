"""A real Office/figure workflow through the shipped CLI in a synthetic root."""

from io import BytesIO
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys

from docx import Document
from PIL import Image
import yaml


SOURCE = Path(__file__).resolve().parents[2]


def invoke(root, *arguments):
    environment = dict(os.environ, PYTHONUTF8="1", PYTHONDONTWRITEBYTECODE="1")
    environment.pop("PYTHONPATH", None)
    result = subprocess.run(
        [sys.executable, "-B", str(root / "knowledge.py"), *arguments, "--json"],
        cwd=root, env=environment, capture_output=True, encoding="utf-8", timeout=180,
    )
    assert result.returncode == 0, (result.stdout, result.stderr)
    return json.loads(result.stdout)


def test_office_figures_p5_cleanup_query_and_independent_source_release(tmp_path):
    project = tmp_path / "project"
    root = project / "himate/synthetic"
    root.mkdir(parents=True)
    for directory in ("runtime", "governance"):
        shutil.copytree(SOURCE / directory, root / directory,
                        ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    for name in ("knowledge.py", "knowledge.sh", "knowledge.cmd"):
        shutil.copyfile(SOURCE / name, root / name)
    (root / "governance/sources.lock.yaml").write_text(
        "schema_version: '2.0'\nresources: []\nsnapshots: []\n", encoding="utf-8")
    (project / "config.yml").write_text(
        "schema_version: '2.0'\nworkspace_id: synthetic\n"
        "knowledge_roots: {wifi: himate/synthetic}\ncode_root: main\nmodules: []\n", encoding="utf-8")
    (root / "docs-raw").mkdir()
    (root / "knowledge").mkdir()
    original = root / "docs-raw/design.docx"
    document = Document()
    document.add_paragraph("The design requires three retry attempts.")
    picture = BytesIO()
    Image.new("RGB", (40, 40), "navy").save(picture, format="PNG")
    picture.seek(0)
    document.add_picture(picture)
    document.save(original)
    original_hash = hashlib.sha256(original.read_bytes()).hexdigest()

    scanned = invoke(root, "ingest", "scan", "--source", "docs-raw/design.docx")["data"]
    source_id = scanned["source"]["source_id"]
    parsed = scanned["source"]["parsed"]
    assert parsed["parse_status"] == "parsed", parsed
    assert parsed["images"] and all((root / item["ref"]).is_file() for item in parsed["images"])
    assert "three retry" in (root / parsed["markdown_ref"]).read_text(encoding="utf-8")
    metadata = {"id": "retry-design", "doc_ref": [
        {"doc_id": source_id, "locator": {"kind": "paragraph", "paragraph": 1}},
    ], "code_ref": []}
    (root / ".build/cards/retry.md").write_text(
        "---\n" + yaml.safe_dump(metadata, sort_keys=False)
        + "---\n# Retry design\n\nretry_count is three by design; runtime behavior is unverified.\n",
        encoding="utf-8")
    (root / ".build/plan.md").write_text(
        "---\ncards: [{id: retry-design, path: knowledge/retry.md}]\n---\n"
        "Preserve the documented retry design.\n", encoding="utf-8")
    assert invoke(root, "ingest", "check")["phase"] == "passed"
    assert invoke(root, "ingest", "commit")["phase"] == "completed"
    assert not (root / ".build").exists()
    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_hash
    assert invoke(root, "search", "retry_count")["data"]["matches"][0]["card_id"] == "retry-design"

    # A later, separate ingestion remains intact while this document is read.
    (root / "docs-raw/next.md").write_text("# Another source\n\nA separate design.\n", encoding="utf-8")
    invoke(root, "ingest", "scan", "--source", "docs-raw/next.md")
    build_before = {p.relative_to(root).as_posix(): p.read_bytes()
                    for p in (root / ".build").rglob("*") if p.is_file()}
    request = root / "artifacts/extract.json"
    request.write_text(json.dumps({"source_id": source_id}), encoding="utf-8")
    reading = invoke(root, "source", "extract", "--input", str(request))["data"]
    assert reading["parsed"]["images"]
    assert all((root / item["ref"]).is_file() for item in reading["parsed"]["images"])
    assert invoke(root, "read", "retry-design")["data"]["card_id"] == "retry-design"
    request.write_text(json.dumps({"read_id": reading["read_id"]}), encoding="utf-8")
    assert invoke(root, "source", "release", "--input", str(request))["data"]["released"] is True
    assert not (root / "artifacts/source-read" / reading["read_id"]).exists()
    assert {p.relative_to(root).as_posix(): p.read_bytes()
            for p in (root / ".build").rglob("*") if p.is_file()} == build_before
