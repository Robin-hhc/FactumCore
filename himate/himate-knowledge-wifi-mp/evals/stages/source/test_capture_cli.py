"""Public boundary tests; no real WiFi originals or human approvals."""

import json
import hashlib
import subprocess
import sys
from pathlib import Path

import pytest


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


def hashed(data):
    return hashlib.sha256(data).hexdigest()


@pytest.mark.parametrize("raw_class", ["specs", "troubleshooting"])
def test_git_source_resolver_recovers_exact_original_without_local_evidence_objects(harness, raw_class):
    import base64
    import yaml

    source_path = "docs-raw/" + raw_class + "/wifi/source.md"
    original_bytes = b"# Source\nCommitted exact version.\n"
    harness.write(source_path, original_bytes)
    captured = harness.run("source", {
        "operation_id": "git-old", "source_path": source_path, "source_id": "synthetic",
    }, action="capture")
    bundle = harness.record(captured["data"]["bundle_ref"])
    workspace = harness.root.parent
    (workspace / "config.yml").write_text(yaml.safe_dump({"schema_version": "2.0", "workspace_id": "synthetic",
        "knowledge_roots": {"wifi": harness.root.name}, "repos": []}), encoding="utf-8")
    def git(*args):
        result = subprocess.run(["git", "-C", str(workspace), *args], capture_output=True, check=True)
        return result.stdout.decode().strip()
    git("init")
    path = harness.root.name + "/" + source_path
    git("add", "--", path)
    git("-c", "user.name=Synthetic", "-c", "user.email=synthetic@example.invalid", "commit", "-m", "Synthetic original only")
    commit = git("rev-parse", "HEAD")
    bundle["snapshots"][0]["recovery"] = {"method": "git", "kind": "git", "commit": commit, "path": path}
    bundle["snapshots"][0].pop("extraction", None)
    bundle.pop("digest")
    bundle["digest"] = hashed(canonical(bundle))
    git_ref = harness.write(".build/operations/git-old/git-capture.json", canonical(bundle))
    original = harness.bundle_path(bundle, "logs/evidence/published/objects/" + bundle["snapshots"][0]["sha256"] + ".bin")
    original.unlink()
    harness.write(source_path, b"# Source\nUncommitted newer version.\n")
    response = harness.run("source", {"operation_id": "git-read", "bundle_ref": git_ref,
        "source_ref": {"evidence_ref": "e1", "source_id": "synthetic", "snapshot_id": bundle["snapshots"][0]["snapshot_id"],
                       "locator": {"kind": "lines", "start": 1, "end": 2}}}, "resolve-source")
    assert base64.b64decode(response["data"]["bytes_b64"]) == original_bytes
    assert response["data"]["current_source_matches"] is False


def test_capture_automatically_registers_matching_committed_raw_git_version(harness):
    import yaml
    workspace = harness.root.parent
    (workspace / "config.yml").write_text(yaml.safe_dump({"knowledge_roots": {"wifi": harness.root.name}, "repos": []}), encoding="utf-8")
    harness.write("docs-raw/specs/source.md", b"# Source\nCommitted bytes.\n")
    def git(*args):
        return subprocess.run(["git", "-C", str(workspace), *args], capture_output=True, check=True).stdout.decode().strip()
    git("init")
    path = harness.root.name + "/docs-raw/specs/source.md"
    git("add", "--", path)
    git("-c", "user.name=Synthetic", "-c", "user.email=synthetic@example.invalid", "commit", "-m", "Synthetic original only")
    response = harness.run("source", {"operation_id": "portable", "source_id": "synthetic", "source_path": "docs-raw/specs/source.md"}, action="capture")
    bundle = harness.record(response["data"]["bundle_ref"])
    snapshot = bundle["snapshots"][0]
    assert snapshot["recovery"] == {"method": "git", "kind": "git", "commit": git("rev-parse", "HEAD"), "path": path, "status": "available"}
    # A second machine has the Git original, not the first machine's evidence cache.
    for row in bundle["staged_objects"]:
        (harness.root / row["stage_ref"]).unlink()
    resolved = harness.run("source", {"operation_id": "portable-read", "bundle_ref": response["data"]["bundle_ref"],
        "source_ref": {"evidence_ref": "e1", "source_id": "synthetic", "snapshot_id": snapshot["snapshot_id"], "locator": {"kind": "lines", "start": 1, "end": 2}}}, "resolve-source")
    assert resolved["data"]["text"] == "# Source\nCommitted bytes.\n"
    checked = harness.run("source", {"operation_id": "portable-check", "bundle_ref": response["data"]["bundle_ref"]}, "resolve")
    assert checked["phase"] == "resolved"


def test_missing_request_returns_structured_error_without_creating_knowledge(tmp_path):
    cli = Path(__file__).resolve().parents[3] / "knowledge.py"
    response = tmp_path / "response.json"
    knowledge = tmp_path / "absent-knowledge"
    result = subprocess.run(
        [
            sys.executable,
            str(cli), "source", "--json",
            "--input",
            str(tmp_path / "missing.json"),
            "--output",
            str(response),
            "--knowledge-root",
            str(knowledge),
            "--rules-digest",
            "0" * 64,
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 2
    record = json.loads(response.read_text(encoding="utf-8") if response.exists() else result.stdout)
    assert record["status"] == "error"
    assert record["diagnostics"]
    assert not knowledge.exists()


@pytest.mark.parametrize(
    "path",
    [
        "../outside.md",
        "docs-raw/sources/link.md",
        "docs-raw/troubleshooting-extra/case.md",
        "docs-raw/troubleshooting/../specs/source.md",
        "https://example.invalid/doc",
        "C:/secret.md",
    ],
)
def test_raw_scope_rejects_outside_and_deferred_inputs(harness, path):
    result = harness.run(
        "source",
        {"operation_id": "unsafe", "source_path": path, "source_id": "synthetic"},
        expected=2,
     action="capture")
    assert result["status"] == "error"
    assert not (harness.root / "logs").exists()


def test_troubleshooting_original_and_attachment_retain_exact_bytes(harness):
    original = "docs-raw/troubleshooting/wifi/case.md"
    attachment = "docs-raw/troubleshooting/wifi/trace.txt"
    inputs = {
        original: b"# Case\nDiagnostic observation: [trace](trace.txt).\n",
        attachment: b"Synthetic observation trace\n",
    }
    for path, data in inputs.items():
        harness.write(path, data)
    result = harness.run("source", {
        "operation_id": "troubleshooting", "source_path": original,
        "source_id": "case", "attachments": [attachment],
    }, action="capture")
    bundle = harness.record(result["data"]["bundle_ref"])
    assert [(resource["path"], resource["kind"]) for resource in bundle["resources"]] == [
        (original, "raw"), (attachment, "attachment"),
    ]
    snapshots = {snapshot["source_id"]: snapshot for snapshot in bundle["snapshots"]}
    for resource in bundle["resources"]:
        path = resource["path"]
        snapshot = snapshots[resource["source_id"]]
        assert snapshot["sha256"] == hashed(inputs[path])
        assert harness.bundle_path(bundle, snapshot["recovery"]["reference"]).read_bytes() == inputs[path]
        assert (harness.root / path).read_bytes() == inputs[path]


def test_exact_source_recovery_survives_original_replacement(harness):
    ref, bundle = harness.capture()
    harness.write("docs-raw/specs/source.md", b"Different newer bytes\n")
    result = harness.run(
        "source", {"operation_id": "op", "bundle_ref": ref}, "resolve"
    )
    snapshot = result["data"]["bundle"]["snapshots"][0]
    assert harness.bundle_path(bundle, snapshot["recovery"]["reference"]).read_bytes() == b"# Source\nSynthetic evidence.\n"
    assert bundle["snapshots"][0] == snapshot


def test_required_attachment_is_not_silently_skipped(harness):
    harness.write("docs-raw/specs/source.md", b"# Doc\n![diagram](diagram.bin)\n")
    harness.write("docs-raw/specs/diagram.bin", b"opaque binary")
    result = harness.run(
        "source",
        {
            "operation_id": "missing",
            "source_path": "docs-raw/specs/source.md",
            "source_id": "synthetic",
        },
        expected=2,
     action="capture")
    assert result["diagnostics"][0]["code"] == "ATTACHMENT_REQUIRED"
    result = harness.run(
        "source",
        {
            "operation_id": "complete",
            "source_path": "docs-raw/specs/source.md",
            "source_id": "synthetic",
            "attachments": ["docs-raw/specs/diagram.bin"],
        },
     action="capture")
    assert len(harness.record(result["data"]["bundle_ref"])["snapshots"]) == 2


def test_retained_byte_tamper_is_detected(harness):
    ref, bundle = harness.capture()
    harness.bundle_path(bundle, bundle["snapshots"][0]["recovery"]["reference"]).write_bytes(
        b"tampered"
    )
    result = harness.run(
        "source",
        {"operation_id": "op", "bundle_ref": ref},
        "resolve",
        expected=2,
    )
    assert result["diagnostics"][0]["code"] == "SOURCE_INTEGRITY_ERROR"


def test_scanned_pdf_does_not_become_success(harness):
    import io
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = io.BytesIO()
    writer.write(stream)
    harness.write("docs-raw/specs/scanned.pdf", stream.getvalue())
    result = harness.run(
        "source",
        {
            "operation_id": "scanned",
            "source_path": "docs-raw/specs/scanned.pdf",
            "source_id": "synthetic",
        },
        expected=3,
     action="capture")
    assert result["diagnostics"][0]["code"] == "PARTIAL_DOCUMENT"


def test_docx_paragraphs_and_tables_have_locators(harness):
    import io
    import docx

    document = docx.Document()
    document.add_paragraph("Synthetic paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic cell"
    stream = io.BytesIO()
    document.save(stream)
    harness.write("docs-raw/specs/doc.docx", stream.getvalue())
    result = harness.run(
        "source",
        {
            "operation_id": "docx",
            "source_path": "docs-raw/specs/doc.docx",
            "source_id": "synthetic",
        },
     action="capture")
    bundle = harness.record(result["data"]["bundle_ref"])
    locators = json.loads(harness.bundle_path(bundle, bundle["extractions"][0]["locator_ref"]).read_bytes())["locators"]
    assert {item["kind"] for item in locators} == {"paragraph", "table_row"}


def test_resolve_source_distinguishes_exact_history_from_current_raw(harness):
    ref, bundle = harness.capture()
    source = {
        "evidence_ref": "e1",
        "source_id": "synthetic",
        "snapshot_id": bundle["snapshots"][0]["snapshot_id"],
        "locator": {"kind": "lines", "start": 2, "end": 2},
    }
    harness.write("docs-raw/specs/source.md", b"A newer raw version\n")
    result = harness.run(
        "source",
        {"operation_id": "op", "bundle_ref": ref, "source_ref": source},
        "resolve-source",
    )
    assert result["data"]["text"] == "Synthetic evidence.\n"
    assert result["data"]["current_source_matches"] is False


def test_source_locator_cannot_select_outside_exact_original(harness):
    ref, bundle = harness.capture()
    source = {
        "evidence_ref": "e1",
        "source_id": "synthetic",
        "snapshot_id": bundle["snapshots"][0]["snapshot_id"],
        "locator": {"kind": "lines", "start": 2, "end": 999},
    }
    result = harness.run(
        "source",
        {"operation_id": "op", "bundle_ref": ref, "source_ref": source},
        "resolve-source",
        expected=2,
    )
    assert result["diagnostics"][0]["code"] == "INVALID_SOURCE_LOCATOR"


def test_unpinned_timestamp_valid_pyc_is_never_executed(harness, tmp_path):
    import importlib.util
    import os
    import py_compile

    target = harness.root / "governance/contracts/validation.py"
    original = target.read_bytes()
    malicious = b"from pathlib import Path\n(Path(__file__).parent / 'PYC_EXECUTED').write_text('bad')\nraise RuntimeError('untrusted cached code')\n"
    attacker = tmp_path / "attacker.py"
    attacker.write_bytes(malicious + b" " * (len(original) - len(malicious)))
    os.utime(attacker, ns=(target.stat().st_atime_ns, target.stat().st_mtime_ns))
    cache = Path(importlib.util.cache_from_source(str(target)))
    cache.parent.mkdir(exist_ok=True)
    py_compile.compile(str(attacker), cfile=str(cache), dfile=str(target), doraise=True)
    harness.capture()
    assert not (target.parent / "PYC_EXECUTED").exists()


def test_text_pdf_has_page_locator(harness):
    import io
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=100)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 50 Td (Synthetic text layer) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    harness.write("docs-raw/specs/text.pdf", stream.getvalue())
    result = harness.run(
        "source",
        {
            "operation_id": "pdf",
            "source_id": "pdf",
            "source_path": "docs-raw/specs/text.pdf",
        },
     action="capture")
    bundle = harness.record(result["data"]["bundle_ref"])
    locator = {
        "evidence_ref": "e1",
        "source_id": "pdf",
        "snapshot_id": bundle["snapshots"][0]["snapshot_id"],
        "locator": {"kind": "page", "page": 1},
    }
    resolved = harness.run(
        "source",
        {
            "operation_id": "pdf",
            "bundle_ref": result["data"]["bundle_ref"],
            "source_ref": locator,
        },
        "resolve-source",
    )
    assert "Synthetic text layer" in resolved["data"]["text"]


def test_payload_output_path_does_not_authorize_host_file_writes(harness, tmp_path):
    victim = tmp_path / "payload-selected.txt"
    harness.write("docs-raw/specs/source.md", b"# Source\nSynthetic evidence.\n")
    harness.run(
        "source",
        {
            "operation_id": "op",
            "source_id": "synthetic",
            "source_path": "docs-raw/specs/source.md",
            "output_path": str(victim),
        },
     action="capture")
    assert not victim.exists()


def test_record_reference_cannot_read_restricted_gold(harness):
    harness.write("evals/gold/secret.json", b"PRIVATE_UNREADABLE_GOLD")
    response = harness.run(
        "source",
        {"operation_id": "op", "bundle_ref": "evals/gold/secret.json"},
        "resolve",
        expected=2,
    )
    assert response["diagnostics"][0]["code"] == "UNSAFE_PATH"
    assert "PRIVATE_UNREADABLE_GOLD" not in str(response)


@pytest.mark.parametrize(
    "workflow",
    [
        {},
        {"run_id": "../escape", "node_id": "capture_raw", "attempt_id": "1"},
        {"run_id": "run", "node_id": "capture_raw"},
    ],
)
def test_workflow_metadata_does_not_gate_source_capture(harness, workflow):
    harness.write("docs-raw/specs/source.md", b"Synthetic evidence.\n")
    response = harness.run(
        "source",
        {
            "operation_id": "op",
            "source_id": "synthetic",
            "source_path": "docs-raw/specs/source.md",
            "workflow": workflow,
        },
     action="capture")
    assert response["status"] == "ok"
    assert (harness.root / response["data"]["bundle_ref"]).is_file()
    assert not (harness.root / "logs/workflows").exists()


def test_source_capture_returns_business_artifact_without_handoff_manifest(harness):
    harness.write("docs-raw/specs/source.md", b"Synthetic evidence.\n")
    response = harness.run(
        "source",
        {
            "operation_id": "capture",
            "source_id": "synthetic",
            "source_path": "docs-raw/specs/source.md",
            "workflow": {
                "run_id": "run",
                "node_id": "capture_raw",
                "attempt_id": "attempt1",
            },
        },
     action="capture")
    reference = response["data"]["bundle_ref"]
    bundle = harness.record(reference)
    assert bundle["digest"] == response["data"]["bundle_digest"]
    assert "artifact_manifest_ref" not in response["data"]
    assert not (harness.root / "logs/workflows").exists()


@pytest.mark.parametrize("module", ["support", "entry"])
def test_unpinned_package_pyc_is_never_executed(harness, tmp_path, module):
    import importlib.util
    import os
    import py_compile

    target = harness.root / "runtime/himate_knowledge/stages/source" / (module + ".py")
    original = target.read_bytes()
    marker = tmp_path / (module + "-PYC-EXECUTED")
    malicious = (
        "open("
        + repr(str(marker))
        + ", 'w').write('bad')\nraise RuntimeError('untrusted package pyc')\n"
    ).encode()
    assert len(malicious) <= len(original)
    attacker = tmp_path / "attacker.py"
    attacker.write_bytes(malicious + b" " * (len(original) - len(malicious)))
    os.utime(attacker, ns=(target.stat().st_atime_ns, target.stat().st_mtime_ns))
    cache = Path(importlib.util.cache_from_source(str(target)))
    cache.parent.mkdir(exist_ok=True)
    prior = cache.read_bytes() if cache.exists() else None
    try:
        py_compile.compile(
            str(attacker), cfile=str(cache), dfile=str(target), doraise=True
        )
        harness.capture()
        assert not marker.exists()
    finally:
        if prior is None:
            cache.unlink(missing_ok=True)
        else:
            cache.write_bytes(prior)
