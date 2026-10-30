"""Real-format dbox journeys against a pinned, separately installed release."""

import hashlib
import base64 as base64
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import time

import pytest
import yaml

from support.releases import load_manual, select_release


RELEASE = "f8a84153a1ce4c23082f70d7de55a2c1f65523f8"
DBOX = "docs-raw/sources/dbox"
CLAIM = "Capture marker is a documented rule."


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def json_bytes(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


@pytest.fixture(scope="session")
def release_runtime(tmp_path_factory):
    checkout = os.environ.get("HIMATE_DBOX_RELEASE_CHECKOUT")
    selection = select_release(Path(__file__).resolve().parents[5], checkout=checkout,
                               manifest_path=os.environ.get("HIMATE_DBOX_RELEASE_MANIFEST"))
    root, manifest = selection.root, selection.manifest
    module = load_manual(selection, "dbox_formats_manual")
    required = selection.knowledge_root / "runtime/manifest.json"
    lock_bytes = (selection.knowledge_root / "governance/framework.lock.yaml").read_bytes()
    identity = {"base_commit": RELEASE, "tool_version": yaml.safe_load(lock_bytes)["tool_version"],
                "rules_sha256": sha(lock_bytes), "revision": manifest.get("revision", "head" if checkout else "current_checkout"),
                "execution_layout": selection.layout,
                "runtime_manifest_sha256": sha(required.read_bytes()) if required.is_file() else None}
    if manifest.get("rules_sha256"):
        assert identity["rules_sha256"] == manifest["rules_sha256"]
    return root, module, identity, selection


@pytest.fixture
def pipeline(tmp_path, release_runtime):
    release, manual, identity, selection = release_runtime

    class AuditedPipeline(manual.ManualPipeline):
        def call(self, slug, request, action=None, expected=0):
            slug, request, action = manual.FIXTURE.target_call(slug, request, action)
            self.sequence += 1
            incoming = self.base / f"format-in-{self.sequence}.json"
            outgoing = self.base / f"format-out-{self.sequence}.json"
            incoming.write_bytes(json_bytes(request))
            if selection.layout == "repository_cli":
                script = self.root / "knowledge.py"
                release_root = selection.knowledge_root
                assert script.read_bytes() == (release_root / "knowledge.py").read_bytes()
                assert (self.root / "runtime/manifest.json").read_bytes() == (release_root / "runtime/manifest.json").read_bytes()
                command = [sys.executable, str(script), slug.removeprefix("knowledge-"), "--json",
                           "--input", str(incoming), "--output", str(outgoing),
                           "--knowledge-root", str(self.root)]
            else:
                command = [*selection.command(self.root, slug, skills=self.skills),
                           "--input", str(incoming), "--output", str(outgoing),
                           "--knowledge-root", str(self.root)]
            if action:
                command += ["--action", action]
            started = time.perf_counter()
            process = subprocess.run(command, capture_output=True, timeout=180)
            elapsed = time.perf_counter() - started
            response = json.loads(outgoing.read_bytes()) if outgoing.exists() else {}
            stdout = self.base / f"format-stdout-{self.sequence}.txt"
            stderr = self.base / f"format-stderr-{self.sequence}.txt"
            stdout.write_bytes(process.stdout)
            stderr.write_bytes(process.stderr)
            self.audit.append({"release": self.release_identity, "skill": "hs-" + slug, "action": action,
                "input": incoming.name, "output": outgoing.name, "returncode": process.returncode,
                "seconds": elapsed, "status": response.get("status"), "phase": response.get("phase"),
                "diagnostics": response.get("diagnostics", []), "stderr": stderr.name})
            (self.base / "cli-audit.json").write_bytes(json_bytes({"release": self.release_identity, "calls": self.audit,
                "seconds": sum(row["seconds"] for row in self.audit)}))
            assert process.returncode == expected, (response, process.stderr.decode(errors="replace"))
            assert response.get("rules_digest") == self.release_identity["rules_sha256"]
            return response

    result = AuditedPipeline(tmp_path)
    result.audit = []
    result.release_identity = identity
    if selection.layout == "repository_cli":
        assert (result.root / "knowledge.py").read_bytes() == (selection.knowledge_root / "knowledge.py").read_bytes()
    assert (result.root / "governance/framework.lock.yaml").read_bytes() == (selection.knowledge_root / "governance/framework.lock.yaml").read_bytes()
    return result


def primary_card(pipeline, bundle, logical, card_id, locator):
    resource = next(row for row in bundle["resources"] if row["path"] == logical)
    actual = next(row for row in bundle["snapshots"] if row["source_id"] == resource["source_id"])
    # Select an actual captured snapshot for card authoring; never modify the
    # immutable bundle's format, extraction identity, bytes or native locators.
    card = pipeline.card(card_id, {"snapshots": [actual]}, text=CLAIM)
    _, front, body = card.decode().split("---", 2)
    metadata = yaml.safe_load(front)
    metadata["sources"][0]["locator"] = locator
    return ("---\n" + yaml.safe_dump(metadata, allow_unicode=True, sort_keys=False) + "---" + body).encode(), metadata


def no_private_path(raw, directory):
    decoded = raw.decode("utf-8", errors="replace")
    values = [decoded]
    serialized = [decoded]
    if decoded.startswith("---") and len(decoded.split("---", 2)) == 3:
        serialized.append(decoded.split("---", 2)[1])
    for text in serialized:
        for parser in (json.loads, yaml.safe_load):
            try:
                pending = [parser(text)]
            except (ValueError, yaml.YAMLError):
                continue
            while pending:
                value = pending.pop()
                if isinstance(value, str):
                    values.append(value)
                elif isinstance(value, dict):
                    pending.extend(value.keys())
                    pending.extend(value.values())
                elif isinstance(value, list):
                    pending.extend(value)
    private = re.sub(r"[\\/]+", "/", str(directory)).casefold()
    assert all(private not in re.sub(r"[\\/]+", "/", value).casefold() for value in values)


def document_bytes(kind):
    if kind == "md":
        return "md", ("# Source\n" + CLAIM + "\n").encode(), {"kind": "lines", "start": 1, "end": 2}
    if kind == "txt":
        return "txt", b"\xef\xbb\xbf# Source\r\n" + CLAIM.encode() + b"\r\n", {"kind": "lines", "start": 1, "end": 2}
    if kind.startswith("docx"):
        import docx
        document = docx.Document()
        document.add_heading("Synthetic format source", level=1)
        if kind == "docx-table":
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = CLAIM
            table.cell(0, 1).text = "Explicit table evidence"
            locator = {"kind": "table_row", "table": 1, "row": 1}
        else:
            document.add_paragraph(CLAIM)
            locator = {"kind": "paragraph", "paragraph": 2}
        stream = io.BytesIO()
        document.save(stream)
        assert stream.getvalue().startswith(b"PK\x03\x04")
        return "docx", stream.getvalue(), locator
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    for text in ("Synthetic context page", CLAIM):
        page = writer.add_blank_page(width=300, height=100)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
            NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(("BT /F1 12 Tf 10 50 Td (" + text + ") Tj ET").encode())
        page[NameObject("/Contents")] = writer._add_object(content)
    stream = io.BytesIO()
    writer.write(stream)
    assert stream.getvalue().startswith(b"%PDF-")
    return "pdf", stream.getvalue(), {"kind": "page", "page": 2}


def invalid_document(kind):
    if kind == "unsupported-csv":
        return "csv", b"column,value\ncapture,marker\n", 2, "UNSUPPORTED_DOCUMENT_FORMAT", "UNSUPPORTED_DOCUMENT_FORMAT"
    if kind == "empty-text":
        return "txt", b"", 2, "UNREADABLE_DOCUMENT", "UNREADABLE_DOCUMENT"
    if kind == "invalid-utf8":
        return "txt", b"\xff\xfe\x00broken UTF8", 2, "UNREADABLE_DOCUMENT", "UNREADABLE_DOCUMENT"
    if kind == "broken-docx":
        return "docx", b"This is not a ZIP document container", 2, "UNREADABLE_DOCUMENT", "UNREADABLE_DOCUMENT"
    if kind == "broken-pdf":
        return "pdf", b"%PDF-1.7\nInvalid synthetic PDF contents\n%%EOF", 2, "UNREADABLE_DOCUMENT", "UNREADABLE_DOCUMENT"
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=100)
    if kind == "encrypted-pdf":
        writer.encrypt("synthetic-test-password")
        code, expected = "ENCRYPTED_DOCUMENT", 2
    else:
        code, expected = "PARTIAL_DOCUMENT", 3
    stream = io.BytesIO()
    writer.write(stream)
    return "pdf", stream.getvalue(), expected, code, code
