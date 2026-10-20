"""Read-only raw acquisition with immutable local recovery objects."""

from __future__ import annotations

import io
import json
import re
from pathlib import PurePosixPath
from urllib.parse import unquote, urlsplit

from support import (
    NodeError,
    canonical,
    digest,
    fixed_bytes,
    load_yaml,
    safe,
    sha,
    write_bytes,
)

RAW_CLASSES = {"conventions", "chips", "storys", "specs", "protocols"}


def extract(data, suffix, *, primary=True):
    if suffix in {".md", ".txt"}:
        text = data.decode("utf-8-sig")
        if "\0" in text or not text.strip():
            raise NodeError(
                "UNREADABLE_DOCUMENT", "UTF-8 text must be nonempty and contain no NUL"
            )
        lines = text.splitlines()
        return (
            text,
            [
                {"kind": "lines", "start": i + 1, "end": i + 1, "text": line}
                for i, line in enumerate(lines)
            ],
            "utf8",
            [],
        )
    if suffix == ".pdf":
        from pypdf import PdfReader, __version__

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise NodeError("ENCRYPTED_DOCUMENT", "Encrypted PDF is unsupported")
        pages = [page.extract_text() or "" for page in reader.pages]
        if not pages or any(not page.strip() for page in pages):
            raise NodeError(
                "PARTIAL_DOCUMENT",
                "PDF page lacks a readable text layer; no OCR",
                blocked=True,
            )
        return (
            "\n\n".join(pages),
            [
                {"kind": "page", "page": i + 1, "text": text}
                for i, text in enumerate(pages)
            ],
            "pypdf:" + __version__,
            ["PDF image content and visual layout are not interpreted"],
        )
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = []
        for i, paragraph in enumerate(document.paragraphs):
            if paragraph.text.strip():
                parts.append(
                    {"kind": "paragraph", "paragraph": i + 1, "text": paragraph.text}
                )
        for i, table in enumerate(document.tables):
            for j, row in enumerate(table.rows):
                parts.append(
                    {
                        "kind": "table_row",
                        "table": i + 1,
                        "row": j + 1,
                        "text": " | ".join(cell.text for cell in row.cells),
                    }
                )
        if not parts or not any(p["text"].strip() for p in parts):
            raise NodeError(
                "PARTIAL_DOCUMENT",
                "DOCX has no readable paragraphs or tables",
                blocked=True,
            )
        if document.element.xpath(".//w:ins | .//w:del"):
            raise NodeError(
                "PARTIAL_DOCUMENT",
                "DOCX tracked changes require explicit resolution",
                blocked=True,
            )
        return (
            "\n\n".join(p["text"] for p in parts),
            parts,
            "python-docx:" + docx.__version__,
            [
                "Embedded images, floating text, headers, and visual table layout are not interpreted"
            ],
        )
    if primary:
        raise NodeError(
            "UNSUPPORTED_DOCUMENT_FORMAT", "Supported primary formats: MD TXT PDF DOCX"
        )
    return (
        None,
        [],
        "opaque",
        ["Attachment bytes retained; content was not interpreted"],
    )


def raw_member(context, ref):
    parts = PurePosixPath(ref).parts if isinstance(ref, str) else ()
    if len(parts) < 3 or parts[0] != "docs-raw" or parts[1] not in RAW_CLASSES:
        raise NodeError(
            "RAW_SCOPE_REQUIRED", "Only five landed docs-raw classes are supported", ref
        )
    path = safe(context.root, ref)
    if not path.is_file():
        raise NodeError("SOURCE_MISSING", "Expected a landed regular raw file", ref)
    before = path.stat()
    data = fixed_bytes(context.root, ref)
    after = path.stat()

    def signature(s):
        return s.st_dev, s.st_ino, s.st_size, s.st_mtime_ns

    if signature(before) != signature(after):
        raise NodeError("SOURCE_CHANGED", "Original changed while reading", ref)
    return {"path": ref, "data": data, "signature": signature(after)}


def check_attachments(members):
    supplied = {member["path"] for member in members}
    gaps = []
    for member in members:
        if PurePosixPath(member["path"]).suffix.lower() != ".md":
            continue
        text = member["data"].decode("utf-8-sig")
        text = re.sub(r"(?ms)^\s*(`{3,}|~{3,}).*?^\s*\1\s*$", "", text)
        text = re.sub(r"(?s)<!--.*?-->|`[^`\n]*`", "", text)
        targets = re.findall(r"!?\[[^\]]*\]\(\s*(<[^>]+>|[^\s)]+)", text)
        targets += re.findall(r"(?im)^\s*\[[^\]]+\]:\s*(<[^>]+>|\S+)", text)
        targets += re.findall(
            r"""(?is)<(?:img|object|iframe|source)\b[^>]*(?:src|data)\s*=\s*["']([^"']+)["']""",
            text,
        )
        targets += re.findall(
            r"""(?is)<(?:img|object|iframe|source)\b[^>]*(?:src|data)\s*=\s*([^\s"'=<>`]+)""",
            text,
        )
        for target in targets:
            target = target.strip("<>")
            parsed = urlsplit(target)
            if parsed.scheme or parsed.netloc or not parsed.path:
                if parsed.scheme or parsed.netloc:
                    gaps.append(
                        "Remote and embedded URI references were not fetched or interpreted"
                    )
                continue  # references remain unvisited; remote bytes never become source evidence.
            path = unquote(parsed.path)
            if path.startswith("/") or "\\" in path or ":" in path:
                raise NodeError(
                    "UNSAFE_ATTACHMENT",
                    "Attachment path must be local raw-relative",
                    path,
                )
            parts = list(PurePosixPath(member["path"]).parent.parts)
            for part in PurePosixPath(path).parts:
                if part == "..":
                    if not parts:
                        raise NodeError(
                            "UNSAFE_ATTACHMENT", "Attachment escapes raw root", target
                        )
                    parts.pop()
                elif part != ".":
                    parts.append(part)
            resolved = "/".join(parts)
            if resolved not in supplied:
                raise NodeError(
                    "ATTACHMENT_REQUIRED",
                    "Referenced local attachment must be explicit",
                    resolved,
                )
    return gaps


def resolve_source(context, reference, bundle=None):
    lock = bundle
    if lock is None:
        lock = load_yaml(fixed_bytes(context.root, "governance/sources.lock.yaml"))
    matches = [
        s
        for s in lock["snapshots"]
        if s["snapshot_id"] == reference["snapshot_id"]
        and s["source_id"] == reference["source_id"]
    ]
    if len(matches) != 1:
        raise NodeError(
            "SOURCE_UNAVAILABLE",
            "Exact source snapshot is not uniquely registered",
            blocked=True,
        )
    snapshot = matches[0]
    recovery = snapshot.get("recovery", {})
    if not isinstance(recovery, dict) or not (
        recovery.get("path") or recovery.get("reference")
    ):
        raise NodeError(
            "RECOVERY_UNAVAILABLE",
            "No supported immutable local recovery object",
            blocked=True,
        )
    if (
        recovery.get("path")
        and recovery.get("reference")
        and recovery["path"] != recovery["reference"]
    ):
        raise NodeError("SOURCE_INTEGRITY_ERROR", "Recovery path aliases disagree")
    original_ref = recovery.get("path", recovery.get("reference"))
    data = fixed_bytes(context.root, original_ref, ["logs/evidence"])
    if sha(data) != snapshot["sha256"] or len(data) != snapshot["bytes"]:
        raise NodeError(
            "SOURCE_INTEGRITY_ERROR", "Recovered bytes differ from exact snapshot"
        )
    resource = next(
        r for r in lock["resources"] if r["source_id"] == reference["source_id"]
    )
    suffix = snapshot.get("format", PurePosixPath(resource["path"]).suffix.lower())
    retained = snapshot.get("extraction")
    if retained:
        extracted = fixed_bytes(context.root, retained["text_ref"], ["logs/evidence"])
        locator_bytes = fixed_bytes(
            context.root, retained["locator_ref"], ["logs/evidence"]
        )
        if (
            sha(extracted) != retained["text_sha256"]
            or sha(locator_bytes) != retained["locator_sha256"]
        ):
            raise NodeError(
                "SOURCE_INTEGRITY_ERROR", "Frozen extraction or locator bytes changed"
            )
        text = extracted.decode("utf-8")
        locators = json.loads(locator_bytes)["locators"]
        gaps = list(retained.get("limitations", []))
    else:
        text, locators, _, gaps = extract(
            data, ".txt" if resource.get("kind") == "code" else suffix, primary=False
        )
        if text is not None:
            gaps.append(
                "This snapshot has no retained extraction; text was re-extracted by the current pinned resolver"
            )
    locator = reference.get("locator", {})
    if text is not None and not locator:
        raise NodeError(
            "INVALID_SOURCE_LOCATOR",
            "Exact source reference requires a readable locator",
        )
    if text is not None and locator:
        kind = locator.get("kind")
        if kind in {"lines", "line"}:
            if resource.get("kind") != "code" and suffix not in {".md", ".txt"}:
                raise NodeError(
                    "INVALID_SOURCE_LOCATOR",
                    "This document format requires its native page/paragraph locator",
                )
            start, end = locator.get("start"), locator.get("end")
            lines = text.splitlines(keepends=True)
            if (
                isinstance(start, bool)
                or isinstance(end, bool)
                or not isinstance(start, int)
                or not isinstance(end, int)
                or not 1 <= start <= end <= len(lines)
            ):
                raise NodeError(
                    "INVALID_SOURCE_LOCATOR",
                    "Source line range is outside exact original",
                )
            text = "".join(lines[start - 1 : end])
        elif kind in {"page", "paragraph", "table_row"}:
            keys = {
                "page": ["page"],
                "paragraph": ["paragraph"],
                "table_row": ["table", "row"],
            }[kind]
            selected = [
                item
                for item in locators
                if item["kind"] == kind
                and all(item.get(key) == locator.get(key) for key in keys)
            ]
            if len(selected) != 1:
                raise NodeError(
                    "INVALID_SOURCE_LOCATOR",
                    "Locator does not identify an extracted original segment",
                )
            text = selected[0]["text"]
        else:
            raise NodeError(
                "UNSUPPORTED_SOURCE_LOCATOR",
                "Unsupported original locator",
                blocked=True,
            )
    try:
        if resource.get("kind") == "code":
            from code_capture import current_code_hash

            current_hash = current_code_hash(context, resource)
        else:
            current_hash = sha(
                fixed_bytes(context.root, resource["path"], ["docs-raw"])
            )
        current_matches = (
            current_hash == snapshot["sha256"] if current_hash is not None else None
        )
    except OSError:
        current_hash, current_matches = None, None
        gaps.append(
            "Current source location is unavailable; only exact retained evidence was recovered"
        )
    if current_matches is None and resource.get("kind") == "code":
        gaps.append(
            "Current registered code root is unavailable or unauthorized; fixed evidence is historical only"
        )
    if current_matches is False:
        gaps.append("Current source bytes differ from this fixed historical snapshot")
    return {
        "source_ref": reference,
        "source_kind": resource.get("kind"),
        "resource": resource,
        "snapshot": snapshot,
        "snapshot_digest": digest(snapshot),
        "original_ref": original_ref,
        "sha256": snapshot["sha256"],
        "bytes": len(data),
        "bytes_b64": __import__("base64").b64encode(data).decode(),
        "text": text,
        "locator": locator,
        "limitations": gaps,
        "current_source_matches": current_matches,
        "current_source_sha256": current_hash,
    }


def execute(context):
    request = context.request
    if context.args.action == "resolve-source":
        bundle = (
            context.record(request["bundle_ref"], "source_bundle")
            if request.get("bundle_ref")
            else None
        )
        return resolve_source(context, request["source_ref"], bundle), "resolved"
    if context.args.action == "resolve":
        bundle = context.record(request["bundle_ref"], "source_bundle")
        for snapshot in bundle["snapshots"]:
            recovery = snapshot["recovery"]
            if (
                recovery.get("path")
                and recovery.get("reference")
                and recovery["path"] != recovery["reference"]
            ):
                raise NodeError(
                    "SOURCE_INTEGRITY_ERROR", "Recovery path aliases disagree"
                )
            data = fixed_bytes(
                context.root,
                recovery.get("path", recovery.get("reference")),
                ["logs/evidence"],
            )
            if sha(data) != snapshot["sha256"] or len(data) != snapshot["bytes"]:
                raise NodeError(
                    "SOURCE_INTEGRITY_ERROR", "Retained object differs from snapshot"
                )
            locator = snapshot.get("locator", {})
            if locator.get("reference") and sha(
                fixed_bytes(context.root, locator["reference"], ["logs/evidence"])
            ) != locator.get("sha256"):
                raise NodeError(
                    "SOURCE_INTEGRITY_ERROR",
                    "Retained locator map differs from snapshot",
                )
        for extraction in bundle.get("extractions", []):
            if (
                sha(
                    fixed_bytes(context.root, extraction["text_ref"], ["logs/evidence"])
                )
                != extraction["text_sha256"]
            ):
                raise NodeError(
                    "SOURCE_INTEGRITY_ERROR", "Retained extraction differs from bundle"
                )
        return {"bundle_ref": request["bundle_ref"], "bundle": bundle}, "resolved"
    if request.get("kind") == "code":
        from code_capture import capture_code

        return capture_code(context)
    if request.get("kind", "raw") != "raw":
        raise NodeError(
            "CODE_CONTEXT_REQUIRED",
            "This node captures landed raw; code needs registered evidence",
            blocked=True,
        )
    source_id = request.get("source_id")
    if (
        not isinstance(source_id, str)
        or not source_id.strip()
        or any(ord(c) < 32 for c in source_id)
    ):
        raise NodeError("INVALID_SOURCE_ID", "Explicit logical source ID required")
    refs = [request["source_path"], *request.get("attachments", [])]
    if len(set(refs)) != len(refs):
        raise NodeError("DUPLICATE_ATTACHMENT", "Each original is enumerated once")
    members = [raw_member(context, ref) for ref in refs]
    reference_gaps = check_attachments(members)
    resources, snapshots, extractions, payloads, limitations = [], [], [], {}, []
    limitations += reference_gaps
    for i, member in enumerate(members):
        ref, data = member["path"], member["data"]
        text, locators, extractor, gaps = extract(
            data, PurePosixPath(ref).suffix.lower(), primary=i == 0
        )
        limitations += gaps
        logical = (
            source_id if i == 0 else source_id + ":attachment:" + sha(ref.encode())[:20]
        )
        object_hash = sha(data)
        snapshot_id = "raw:" + digest(
            {
                "source_id": logical,
                "sha256": object_hash,
                "version": request.get("version"),
                "extractor": extractor,
                "text_digest": sha(text.encode()) if text is not None else None,
                "locator_digest": digest(locators),
            }
        )
        object_ref = "logs/evidence/objects/" + object_hash + ".bin"
        locator_ref = "logs/evidence/locators/" + digest(locators) + ".json"
        payloads[object_ref], payloads[locator_ref] = (
            data,
            canonical({"locators": locators}),
        )
        resources.append(
            {
                "source_id": logical,
                "kind": "raw" if i == 0 else "attachment",
                "path": ref,
            }
        )
        snapshots.append(
            {
                "source_id": logical,
                "snapshot_id": snapshot_id,
                "sha256": object_hash,
                "bytes": len(data),
                "version": request.get("version"),
                "format": PurePosixPath(ref).suffix.lower(),
                "locator": {
                    "kind": "locator_map",
                    "reference": locator_ref,
                    "sha256": sha(payloads[locator_ref]),
                },
                "recovery": {
                    "method": "evidence_object",
                    "path": object_ref,
                    "reference": object_ref,
                    "status": "available",
                },
            }
        )
        if text is not None:
            text_ref = "logs/evidence/extractions/" + sha(text.encode()) + ".txt"
            payloads[text_ref] = text.encode()
            extractions.append(
                {
                    "snapshot_id": snapshot_id,
                    "extractor": extractor,
                    "text_ref": text_ref,
                    "text_sha256": sha(text.encode()),
                    "locator_ref": locator_ref,
                    "locator_sha256": sha(payloads[locator_ref]),
                    "limitations": gaps,
                }
            )
            snapshots[-1]["extraction"] = dict(extractions[-1])
    for member in members:
        now = raw_member(context, member["path"])
        if now["data"] != member["data"] or now["signature"] != member["signature"]:
            raise NodeError(
                "SOURCE_CHANGED",
                "Original changed before evidence publication",
                member["path"],
            )
    bundle = context.seal(
        {
            "resources": resources,
            "snapshots": snapshots,
            "extractions": extractions,
            "context_ref": request.get("context_ref"),
            "limitations": sorted(set(limitations)),
        },
        "source_bundle",
    )
    ref = context.base + "/capture.json"
    existing = safe(context.root, ref)
    if existing.exists() and existing.read_bytes() != canonical(bundle):
        raise NodeError(
            "OPERATION_CONFLICT", "Operation already captures different originals", ref
        )
    for path, data in payloads.items():
        write_bytes(context.root, path, data)
    for member in members:
        now = raw_member(context, member["path"])
        if now["data"] != member["data"] or now["signature"] != member["signature"]:
            raise NodeError(
                "SOURCE_CHANGED",
                "Original changed during evidence publication",
                member["path"],
            )
    context.save(ref, bundle)
    lock_path = safe(context.root, "governance/sources.lock.yaml")
    lock = (
        load_yaml(lock_path.read_bytes())
        if lock_path.exists()
        else {"resources": [], "snapshots": []}
    )
    changes = []
    for snapshot in snapshots:
        prior = [
            s
            for s in lock.get("snapshots", [])
            if s["source_id"] == snapshot["source_id"]
        ]
        if prior and snapshot["snapshot_id"] not in {s["snapshot_id"] for s in prior}:
            changes.append(
                {
                    "source_id": snapshot["source_id"],
                    "old_snapshot_ids": sorted(s["snapshot_id"] for s in prior),
                    "new_snapshot_id": snapshot["snapshot_id"],
                }
            )
    if changes:
        response = context.call(
            "hs-knowledge-maintain",
            {"operation_id": context.operation, "bundle_ref": ref, "changes": changes},
            "register-source-change",
        )
        event_refs = response["data"].get("event_refs", [])
    else:
        event_refs = []
    return {
        "bundle_ref": ref,
        "bundle_digest": bundle["digest"],
        "event_refs": event_refs,
        "limitations": bundle["limitations"],
    }, "captured"
